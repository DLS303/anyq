"""Marketplace reads and writes. One database transaction for selecting an offer."""

from __future__ import annotations

import csv
import io
import time
from datetime import date, datetime, timedelta, timezone
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import delete, func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.marketplace_bridge import on_offer_selected, on_request_sent
from app.marketplace_models import (
    CatalogItem,
    Deal,
    MarketMember,
    MarketRequest,
    MarketRequestLine,
    Offer,
    OfferLine,
    PriceIndexDaily,
    PriceItem,
    PriceObservation,
    SupplierProfile,
    SupplierRating,
    SupplierReview,
)
from app.marketplace_price import (
    HAS_MEDIAN,
    LOW_DATA,
    MIN_SUPPLIERS,
    SOURCE_TYPES,
    WINDOW_DAYS,
    PricePoint,
    excl_tiyn,
    has_source,
    incl_tiyn,
    in_window,
    judge_price,
    market_label,
    market_today,
    overpay_parts,
    summarize,
)
from app.marketplace_score import (
    NEUTRAL_RATING,
    DealFact,
    Score,
    aggregate,
)
from app.sql_backend import RequestRecipient

REVIEW_LIMIT = 5
REVIEW_WINDOW_S = 3600
EDIT_DAYS = 14
_review_hits: dict[str, list[float]] = {}

CATEGORIES = ("Металлопрокат", "Сэндвич-панели", "Крепёж", "Утеплитель")
VERIFICATION_LABELS = {
    "not_verified": "не проверен",
    "verified_ai": "проверен ИИ",
    "verified_procurement": "проверен снабжением",
}


class MarketError(Exception):
    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def _now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def _today() -> date:
    return market_today()


def reset_review_limit() -> None:
    _review_hits.clear()


def _hit_review_limit(user_id: str) -> None:
    now = time.time()
    hits = [stamp for stamp in _review_hits.get(user_id, []) if now - stamp < REVIEW_WINDOW_S]
    if len(hits) >= REVIEW_LIMIT:
        raise MarketError("Слишком много оценок. Попробуйте позже", 429)
    hits.append(now)
    _review_hits[user_id] = hits


def e164(value: str | None) -> str:
    """Keep a phone as + and digits so uniqueness does not depend on spacing."""

    raw = str(value or "").strip()
    digits = "".join(ch for ch in raw if ch.isdigit())
    if not digits:
        return ""
    if len(digits) == 11 and digits.startswith("8"):
        digits = "7" + digits[1:]
    return "+" + digits


def search_text_of(name: str, category: str = "", brand: str = "") -> str:
    return " ".join(part.casefold() for part in (name, category, brand) if part).strip()


def _like(value: str) -> str:
    return value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def stale(item: CatalogItem, today: date | None = None) -> bool:
    today = today or _today()
    try:
        start = date.fromisoformat(item.price_date[:10])
    except ValueError:
        return True
    return today > start + timedelta(days=int(item.valid_days or 0))


def availability(item: CatalogItem) -> str:
    if item.stock == "in_stock" and int(item.lead_days or 0) <= 0:
        return "в наличии"
    days = int(item.lead_days or 0)
    if days > 0:
        return f"{days} дней"
    return "под заказ"


def vat_label(vat_included: bool, vat_rate: int) -> str:
    if not vat_included or vat_rate <= 0:
        return "без НДС"
    return "с НДС"


def _member(session: Session, user_id: str, kind: str | None = None) -> MarketMember | None:
    stmt = select(MarketMember).where(MarketMember.user_id == user_id)
    if kind:
        stmt = stmt.where(MarketMember.kind == kind)
    return session.scalars(stmt).first()


def company_of(session: Session, user, kind: str) -> str:
    row = _member(session, user.id, kind)
    if row:
        return row.company_id
    if kind == "supplier" and user.role == "supplier":
        other = _member(session, user.id, None)
        if other:
            return other.company_id
    from app.store import company_key

    return company_key(user)


def _profile(session: Session, company_id: str) -> SupplierProfile:
    row = session.get(SupplierProfile, company_id)
    if row is None or row.hidden:
        raise MarketError("Поставщик не найден", 404)
    return row


def _rating_row(session: Session, company_id: str) -> SupplierRating | None:
    return session.get(SupplierRating, company_id)


def _facts(session: Session, company_id: str) -> list[DealFact]:
    deals = session.scalars(
        select(Deal).where(Deal.supplier_company_id == company_id, Deal.accepted_at.is_not(None))
    ).all()
    reviews = {
        row.deal_id: row
        for row in session.scalars(
            select(SupplierReview).where(SupplierReview.supplier_company_id == company_id)
        ).all()
    }
    facts: list[DealFact] = []
    for deal in deals:
        review = reviews.get(deal.id)
        try:
            closed = date.fromisoformat((deal.accepted_at or "")[:10])
        except ValueError:
            continue
        facts.append(
            DealFact(
                closed_on=closed,
                on_time=bool(deal.on_time),
                price_matched=bool(deal.price_matched),
                acceptance_ok=True if deal.acceptance_ok is None else bool(deal.acceptance_ok),
                esf_ok=bool(deal.esf_ok),
                response_hours=float(deal.response_hours or 24),
                quality_star=float(review.quality) if review else None,
            )
        )
    return facts


def recompute_supplier(session: Session, company_id: str, as_of: date | None = None) -> Score:
    as_of = as_of or _today()
    score = aggregate(_facts(session, company_id), as_of)
    reviews = session.scalar(
        select(func.count()).select_from(SupplierReview).where(SupplierReview.supplier_company_id == company_id)
    )
    now = _now()
    row = session.get(SupplierRating, company_id)
    if row is None:
        row = SupplierRating(company_id=company_id, rating=score.rating, updated_at=now)
        session.add(row)
    row.rating = score.rating
    row.deal_count = score.deal_count
    row.review_count = int(reviews or 0)
    row.is_new = score.is_new
    row.reliable = score.reliable
    row.delivery_score = score.delivery
    row.price_score = score.price
    row.quality_score = score.quality
    row.documents_score = score.documents
    row.reply_score = score.reply
    row.updated_at = now
    profile = session.get(SupplierProfile, company_id)
    if profile is not None:
        profile.rating = score.rating
        profile.deal_count = score.deal_count
        profile.reliable = score.reliable
    return score


def recompute_all(session: Session, as_of: date | None = None) -> int:
    ids = list(session.scalars(select(SupplierProfile.company_id)).all())
    for company_id in ids:
        recompute_supplier(session, company_id, as_of)
    return len(ids)


def recompute_price_index(session: Session, as_of: date | None = None) -> int:
    """Replace one day's index from append-only observations.

    Suspicious prices are left out. A city with fewer than three distinct
    suppliers keeps its counts and an empty median.
    """

    as_of = as_of or _today()
    day = as_of.isoformat()
    grouped: dict[tuple[str, str], list[PricePoint]] = {}
    for row in session.scalars(select(PriceObservation)).all():
        if row.source_type not in SOURCE_TYPES or not has_source(row.source_type, row.source_id):
            continue
        if not in_window(row.observed_on, as_of):
            continue
        grouped.setdefault((row.item_id, row.city or ""), []).append(
            PricePoint(row.company_id, row.price_tiyn, row.vat_included, row.vat_rate)
        )
    session.execute(delete(PriceIndexDaily).where(PriceIndexDaily.day == day))
    for (item_id, city), points in grouped.items():
        snap = summarize(points)
        session.add(
            PriceIndexDaily(
                item_id=item_id,
                city=city,
                day=day,
                median_tiyn=snap.median_tiyn,
                p25_tiyn=snap.p25_tiyn,
                p75_tiyn=snap.p75_tiyn,
                n_observations=snap.n_observations,
                n_suppliers=snap.n_suppliers,
            )
        )
    session.flush()
    return len(grouped)


def _catalog_index(session: Session) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for item in session.scalars(select(PriceItem)).all():
        for catalog_id in item.catalog_ids or []:
            mapping[str(catalog_id)] = item.id
    return mapping


def _observations(session: Session, item_ids: set[str]) -> list[PriceObservation]:
    if not item_ids:
        return []
    return list(
        session.scalars(select(PriceObservation).where(PriceObservation.item_id.in_(item_ids))).all()
    )


def _empty_market() -> dict:
    return {
        "median_tiyn": None,
        "median_label": LOW_DATA,
        "market_city": "",
        "market_label": "",
        "points": 0,
        "n_suppliers": 0,
        "p25_tiyn": None,
        "p75_tiyn": None,
    }


def _index_row(session: Session, item_id: str, city: str, as_of: date) -> PriceIndexDaily | None:
    start = (as_of - timedelta(days=WINDOW_DAYS)).isoformat()
    return session.scalar(
        select(PriceIndexDaily)
        .where(
            PriceIndexDaily.item_id == item_id,
            PriceIndexDaily.city == city,
            PriceIndexDaily.day >= start,
            PriceIndexDaily.day <= as_of.isoformat(),
        )
        .order_by(PriceIndexDaily.day.desc())
        .limit(1)
    )


def _market_from_row(row: PriceIndexDaily | None) -> dict:
    if row is None or row.median_tiyn is None or row.n_suppliers < MIN_SUPPLIERS:
        return _empty_market()
    return {
        "median_tiyn": row.median_tiyn,
        "median_label": HAS_MEDIAN,
        "market_city": row.city,
        "market_label": market_label(row.city),
        "points": row.n_observations,
        "n_suppliers": row.n_suppliers,
        "p25_tiyn": row.p25_tiyn,
        "p75_tiyn": row.p75_tiyn,
    }


def medians_by_catalog(session: Session, catalog_ids: list[str], as_of: date | None = None) -> dict[str, dict]:
    as_of = as_of or _today()
    mapping = _catalog_index(session)
    out: dict[str, dict] = {}
    for catalog_id in catalog_ids:
        item = session.get(CatalogItem, catalog_id)
        item_id = mapping.get(catalog_id)
        city = item.city if item is not None else ""
        row = _index_row(session, item_id, city, as_of) if item_id else None
        out[catalog_id] = _market_from_row(row)
    return out


def _last_confirmed(session: Session, item_id: str, city: str) -> int | None:
    last: tuple[str, int] | None = None
    for row in _observations(session, {item_id}):
        if row.city != city or not has_source(row.source_type, row.source_id):
            continue
        value = incl_tiyn(row.price_tiyn, row.vat_included, row.vat_rate)
        if last is None or row.observed_on > last[0]:
            last = (row.observed_on, value)
    return last[1] if last else None


def _judge_line(session: Session, catalog_id: str | None, name: str, price_incl: int, as_of: date) -> dict:
    mapping = _catalog_index(session)
    item_id = mapping.get(catalog_id or "")
    catalog = session.get(CatalogItem, catalog_id) if catalog_id else None
    city = catalog.city if catalog is not None else ""
    row = _index_row(session, item_id, city, as_of) if item_id else None
    last = _last_confirmed(session, item_id, city) if item_id else None
    judged = judge_price(
        price_incl,
        [],
        last,
        index_median=row.median_tiyn if row is not None else None,
        index_suppliers=row.n_suppliers if row is not None else 0,
    )
    judged["name"] = name
    return judged


def stars_of(rating: SupplierRating | None, profile: SupplierProfile | None = None) -> dict:
    value = NEUTRAL_RATING
    count = 0
    is_new = True
    reliable = False
    if rating is not None:
        value = float(rating.rating)
        count = int(rating.deal_count or 0)
        is_new = bool(rating.is_new)
        reliable = bool(rating.reliable)
    elif profile is not None and profile.rating is not None:
        value = float(profile.rating)
        count = int(profile.deal_count or 0)
        is_new = count < 5
        reliable = bool(profile.reliable)
    return {
        "rating": round(value, 2),
        "deal_count": count,
        "review_count": int(rating.review_count) if rating else 0,
        "is_new": is_new,
        "label": "Новый" if is_new else "",
        "badge": "Надёжный поставщик" if reliable else "",
        "reliable": reliable,
    }


def card_from(item: CatalogItem, profile: SupplierProfile, rating: SupplierRating | None, median: dict) -> dict:
    expired = stale(item)
    shown = incl_tiyn(item.price_tiyn, item.vat_included, item.vat_rate)
    overpay_pct, overpay_label = overpay_parts(shown, median.get("median_tiyn"))
    return {
        "id": item.id,
        "name": item.name,
        "brand": item.brand,
        "category": item.category,
        "unit": item.unit,
        "price_tiyn": item.price_tiyn,
        "vat_included": item.vat_included,
        "vat_rate": item.vat_rate,
        "vat_label": vat_label(item.vat_included, item.vat_rate),
        "stale": expired,
        "stale_label": "цена устарела" if expired else "",
        "availability": availability(item),
        "city": item.city or profile.city,
        "stock": item.stock,
        "lead_days": item.lead_days,
        "supplier_id": profile.company_id,
        "supplier_name": profile.name,
        "price_date": item.price_date,
        "valid_days": item.valid_days,
        "source_kind": item.source_kind,
        "hidden": item.hidden,
        "median_tiyn": median.get("median_tiyn"),
        "median_label": median.get("median_label") or LOW_DATA,
        "market_city": median.get("market_city") or "",
        "market_label": median.get("market_label") or "",
        "overpay_pct": overpay_pct,
        "overpay_label": overpay_label,
        **stars_of(rating, profile),
    }


def _cards(session: Session, rows: list[tuple[CatalogItem, SupplierProfile, SupplierRating | None]]) -> list[dict]:
    medians = medians_by_catalog(session, [item.id for item, _profile, _rating in rows])
    return [card_from(item, profile, rating, medians.get(item.id, {})) for item, profile, rating in rows]


def home() -> dict:
    return {
        "categories": [{"id": name, "title": name} for name in CATEGORIES],
        "search": True,
    }


def search_catalog(
    session: Session,
    q: str = "",
    category: str = "",
    city: str = "",
    price_min: int | None = None,
    price_max: int | None = None,
    limit: int = 20,
    offset: int = 0,
) -> tuple[list[dict], int]:
    needle = q.casefold().strip()
    filters = [CatalogItem.hidden.is_(False), SupplierProfile.hidden.is_(False)]
    if category:
        filters.append(CatalogItem.category == category)
    if city:
        filters.append(or_(CatalogItem.city == city, SupplierProfile.city == city))
    if price_min is not None:
        filters.append(CatalogItem.price_tiyn >= price_min)
    if price_max is not None:
        filters.append(CatalogItem.price_tiyn <= price_max)
    if needle:
        safe = f"%{_like(needle)}%"
        dialect = session.get_bind().dialect.name
        if dialect == "postgresql":
            filters.append(CatalogItem.name.ilike(safe, escape="\\"))
        else:
            filters.append(CatalogItem.search_text.like(safe, escape="\\"))
    joined = (
        select(CatalogItem, SupplierProfile, SupplierRating)
        .join(SupplierProfile, CatalogItem.company_id == SupplierProfile.company_id)
        .outerjoin(SupplierRating, SupplierRating.company_id == SupplierProfile.company_id)
        .where(*filters)
    )
    total = int(
        session.scalar(
            select(func.count())
            .select_from(CatalogItem)
            .join(SupplierProfile, CatalogItem.company_id == SupplierProfile.company_id)
            .where(*filters)
        )
        or 0
    )
    relevance_exact = func.lower(CatalogItem.name) == needle if needle else CatalogItem.id.is_not(None)
    # Relevance first, then a blend of the cached rating and a lower price.
    # New suppliers already carry the neutral score. No paid term is added.
    rating_value = func.coalesce(SupplierRating.rating, NEUTRAL_RATING)
    price_term = 1.0 / (1.0 + CatalogItem.price_tiyn / 1_000_000.0)
    blend = (0.55 * (rating_value / 5.0)) + (0.45 * price_term)
    if needle:
        order = [relevance_exact.desc(), blend.desc(), CatalogItem.price_tiyn.asc(), CatalogItem.id.asc()]
    else:
        order = [blend.desc(), CatalogItem.price_tiyn.asc(), CatalogItem.id.asc()]
    rows = session.execute(joined.order_by(*order).limit(limit).offset(offset)).all()
    packed = [(item, profile, rating) for item, profile, rating in rows]
    return _cards(session, packed), total


def cities(session: Session) -> list[str]:
    values = session.scalars(
        select(CatalogItem.city)
        .where(CatalogItem.hidden.is_(False), CatalogItem.city != "")
        .distinct()
        .order_by(CatalogItem.city)
    ).all()
    return list(values)


def supplier_public(session: Session, company_id: str, limit: int = 20, offset: int = 0) -> dict:
    profile = _profile(session, company_id)
    rating = _rating_row(session, company_id)
    filters = [CatalogItem.company_id == company_id, CatalogItem.hidden.is_(False)]
    total = int(
        session.scalar(select(func.count()).select_from(CatalogItem).where(*filters)) or 0
    )
    items = session.scalars(
        select(CatalogItem).where(*filters).order_by(CatalogItem.name).limit(limit).offset(offset)
    ).all()
    cards = _cards(session, [(item, profile, rating) for item in items])
    reviews = session.scalars(
        select(SupplierReview)
        .where(SupplierReview.supplier_company_id == company_id)
        .order_by(SupplierReview.created_at.desc())
        .limit(8)
    ).all()
    return {
        "id": profile.company_id,
        "name": profile.name,
        "city": profile.city,
        "vat_payer": profile.vat_payer,
        "vat_rate": profile.vat_rate,
        "vat_label": "без НДС" if not profile.vat_payer or profile.vat_rate <= 0 else "с НДС",
        "esf": profile.esf,
        "esf_label": "ЭСФ" if profile.esf else "без ЭСФ",
        "categories": list(profile.categories or []),
        "delivery": profile.delivery,
        "delivery_cities": list(profile.delivery_cities or []),
        "min_order_tiyn": profile.min_order_tiyn,
        "verification": profile.verification,
        "verification_label": VERIFICATION_LABELS.get(profile.verification, "не проверен"),
        **stars_of(rating, profile),
        "items": cards,
        "total": total,
        "reviews": [
            {
                "id": row.id,
                "quality": row.quality,
                "delivery": row.delivery,
                "price": row.price,
                "communication": row.communication,
                "comment": row.comment,
                "supplier_reply": row.supplier_reply or "",
                "created_at": row.created_at,
            }
            for row in reviews
        ],
    }


def _require_source(source: str) -> str:
    text = (source or "").strip()
    if not text:
        raise MarketError("Укажите источник цены")
    return text


def _own_profile(session: Session, user) -> SupplierProfile:
    if user.role not in ("supplier", "admin"):
        raise MarketError("Недостаточно прав", 403)
    company_id = company_of(session, user, "supplier")
    profile = session.get(SupplierProfile, company_id)
    if profile is None:
        raise MarketError("Нет компании поставщика", 403)
    return profile


def list_own(session: Session, user) -> list[dict]:
    profile = _own_profile(session, user)
    items = session.scalars(
        select(CatalogItem).where(CatalogItem.company_id == profile.company_id).order_by(CatalogItem.name)
    ).all()
    rating = _rating_row(session, profile.company_id)
    return [
        card_from(item, profile, rating, {"median_label": LOW_DATA, "median_tiyn": None})
        | {"hidden": item.hidden, "source": item.source}
        for item in items
    ]


def save_item(session: Session, user, body: dict, item_id: str | None = None) -> dict:
    profile = _own_profile(session, user)
    source = _require_source(str(body.get("source") or ""))
    name = " ".join(str(body.get("name") or "").split())
    if not name:
        raise MarketError("Укажите название")
    category = str(body.get("category") or "").strip()
    if category not in CATEGORIES:
        raise MarketError("Выберите категорию")
    unit = str(body.get("unit") or "").strip()
    if unit not in {"м", "м²", "м2", "т", "шт", "уп"}:
        raise MarketError("Единица: м, м², т, шт или уп")
    if unit == "м2":
        unit = "м²"
    try:
        price = int(body.get("price_tiyn"))
    except (TypeError, ValueError) as exc:
        raise MarketError("Цена в тиынах, целое число") from exc
    if price <= 0:
        raise MarketError("Цена в тиынах, целое число")
    vat_rate = int(body.get("vat_rate") if body.get("vat_rate") is not None else (16 if profile.vat_payer else 0))
    if vat_rate not in (0, 16):
        raise MarketError("Ставка НДС 16% или 0")
    vat_included = bool(body.get("vat_included")) if body.get("vat_included") is not None else profile.vat_payer
    if not profile.vat_payer:
        vat_included = False
        vat_rate = 0
    brand = str(body.get("brand") or "")[:200]
    sku = str(body.get("sku") or "").strip()[:64]
    now = _now()
    if item_id:
        item = session.get(CatalogItem, item_id)
        if item is None or item.company_id != profile.company_id:
            raise MarketError("Позиция не найдена", 404)
    else:
        item = CatalogItem(
            id=uuid4().hex[:16],
            company_id=profile.company_id,
            created_at=now,
            price_tiyn=price,
            price_date=_today().isoformat(),
            source=source,
            name=name,
            category=category,
            unit=unit,
        )
        session.add(item)
    item.name = name
    item.brand = brand
    if sku:
        with session.no_autoflush:
            taken = session.scalar(
                select(CatalogItem.id).where(
                    CatalogItem.company_id == profile.company_id,
                    CatalogItem.sku == sku,
                    CatalogItem.id != item.id,
                )
            )
        if taken:
            raise MarketError("Такой артикул уже есть", 409)
    item.sku = sku
    item.category = category
    item.unit = unit
    item.price_tiyn = price
    item.vat_included = vat_included
    item.vat_rate = vat_rate
    item.price_date = str(body.get("price_date") or item.price_date or _today().isoformat())[:20]
    item.valid_days = int(body.get("valid_days") or 7)
    item.source = source
    item.source_kind = str(body.get("source_kind") or "file")[:20]
    item.stock = "in_stock" if str(body.get("stock") or "in_stock") == "in_stock" else "to_order"
    item.lead_days = max(0, int(body.get("lead_days") or 0))
    item.city = str(body.get("city") or profile.city)[:80]
    item.hidden = bool(body.get("hidden", item.hidden if item_id else False))
    item.search_text = search_text_of(item.name, item.category, item.brand)
    item.updated_at = now
    session.flush()
    return card_from(item, profile, _rating_row(session, profile.company_id), {"median_label": LOW_DATA})


def hide_item(session: Session, user, item_id: str, hidden: bool) -> dict:
    profile = _own_profile(session, user)
    item = session.get(CatalogItem, item_id)
    if item is None or item.company_id != profile.company_id:
        raise MarketError("Позиция не найдена", 404)
    item.hidden = hidden
    item.updated_at = _now()
    return {"id": item.id, "hidden": item.hidden}


def _rows_from_csv(raw: bytes) -> list[dict]:
    text = raw.decode("utf-8-sig", errors="replace")
    sample = text.splitlines()[0] if text else ""
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    return list(csv.DictReader(io.StringIO(text), delimiter=delimiter))


def _rows_from_xlsx(raw: bytes) -> list[dict]:
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = book.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(cell or "").strip() for cell in rows[0]]
    out = []
    for row in rows[1:]:
        if not any(cell not in (None, "") for cell in row):
            continue
        out.append({header[index]: row[index] for index in range(min(len(header), len(row)))})
    return out


_COLUMNS = {
    "name": ("name", "название", "наименование"),
    "category": ("category", "категория"),
    "unit": ("unit", "единица", "ед"),
    "price_tiyn": ("price_tiyn", "цена_тиын", "цена"),
    "vat_included": ("vat_included", "ндс"),
    "source": ("source", "источник"),
    "city": ("city", "город"),
    "stock": ("stock", "наличие"),
    "lead_days": ("lead_days", "срок_дней", "срок"),
    "brand": ("brand", "марка", "гост"),
    "sku": ("sku", "артикул", "код"),
}


def _cell(row: dict, key: str) -> str:
    wanted = _COLUMNS[key]
    lowered = {str(name).strip().casefold(): value for name, value in row.items()}
    for name in wanted:
        if name in lowered and lowered[name] not in (None, ""):
            return str(lowered[name]).strip()
    return ""


def import_prices(session: Session, user, raw: bytes, filename: str) -> dict:
    profile = _own_profile(session, user)
    name = filename.lower()
    if name.endswith(".xlsx"):
        rows = _rows_from_xlsx(raw)
    else:
        rows = _rows_from_csv(raw)
    if not rows:
        raise MarketError("В файле нет строк")
    saved = []
    for index, row in enumerate(rows, start=2):
        source = _cell(row, "source")
        if not source:
            raise MarketError(f"Строка {index}: нет источника цены")
        price_text = _cell(row, "price_tiyn").replace(" ", "").replace(",", ".")
        if "." in price_text:
            raise MarketError(f"Строка {index}: цена в тиынах, без дробной части")
        vat_cell = _cell(row, "vat_included").casefold()
        vat_included = vat_cell not in {"0", "нет", "false", "без"}
        saved.append(
            save_item(
                session,
                user,
                {
                    "name": _cell(row, "name"),
                    "category": _cell(row, "category"),
                    "unit": _cell(row, "unit"),
                    "price_tiyn": price_text,
                    "vat_included": vat_included,
                    "source": source,
                    "city": _cell(row, "city") or profile.city,
                    "stock": "to_order" if _cell(row, "stock").casefold() in {"под заказ", "to_order"} else "in_stock",
                    "lead_days": _cell(row, "lead_days") or 0,
                    "brand": _cell(row, "brand"),
                    "sku": _cell(row, "sku"),
                    "source_kind": "file",
                },
            )
        )
    return {"saved": len(saved), "items": saved}


def my_rating(session: Session, user) -> dict:
    profile = _own_profile(session, user)
    rating = _rating_row(session, profile.company_id)
    stars = stars_of(rating, profile)
    tips = _tips(rating, stars["is_new"])
    return {**stars, "tips": tips, "name": profile.name}


def _tips(rating: SupplierRating | None, is_new: bool) -> list[str]:
    if is_new or rating is None:
        return ["После пяти завершённых сделок покупатели увидят оценку. Сейчас стоит метка «Новый»."]
    parts = {
        "delivery": (rating.delivery_score, "Привозите в обещанный срок — это сильнее всего влияет на оценку."),
        "price": (rating.price_score, "Держите цену счёта равной цене в предложении."),
        "quality": (rating.quality_score, "Сдавайте материал без замечаний по качеству."),
        "documents": (rating.documents_score, "Прикладывайте счёт и ЭСФ без ошибок."),
        "reply": (rating.reply_score, "Отвечайте на заявку в течение суток."),
    }
    weakest = min(parts.values(), key=lambda item: item[0])
    return [weakest[1]]


def _lines_from_body(session: Session, payload: list[dict]) -> list[MarketRequestLine]:
    lines: list[MarketRequestLine] = []
    for raw in payload:
        catalog_id = str(raw.get("catalog_id") or "") or None
        item = session.get(CatalogItem, catalog_id) if catalog_id else None
        name = " ".join(str(raw.get("name") or (item.name if item else "")).split())
        if not name:
            raise MarketError("Укажите позицию")
        try:
            qty = float(raw.get("qty") or 1)
        except (TypeError, ValueError) as exc:
            raise MarketError("Количество больше нуля") from exc
        if qty <= 0:
            raise MarketError("Количество больше нуля")
        lines.append(
            MarketRequestLine(
                id=uuid4().hex[:16],
                request_id="",
                catalog_id=item.id if item else None,
                name=name,
                unit=str(raw.get("unit") or (item.unit if item else "шт")),
                qty=qty,
                category=item.category if item else str(raw.get("category") or ""),
            )
        )
    if not lines:
        raise MarketError("Добавьте хотя бы одну позицию")
    return lines


def _suppliers_for(session: Session, lines: list[MarketRequestLine], explicit: list[str] | None) -> list[str]:
    if explicit:
        return list(dict.fromkeys(explicit))
    ids: list[str] = []
    for line in lines:
        if not line.catalog_id:
            continue
        item = session.get(CatalogItem, line.catalog_id)
        if item and item.company_id not in ids:
            ids.append(item.company_id)
    if not ids:
        raise MarketError("Выберите поставщика")
    return ids


def create_request(session: Session, user, body: dict) -> dict:
    buyer = company_of(session, user, "buyer")
    lines = _lines_from_body(session, list(body.get("lines") or []))
    suppliers = _suppliers_for(session, lines, list(body.get("supplier_ids") or []) or None)
    for company_id in suppliers:
        if session.get(SupplierProfile, company_id) is None:
            raise MarketError("Поставщик не найден", 404)
    now = _now()
    request = MarketRequest(
        id=uuid4().hex[:16],
        buyer_company_id=buyer,
        created_by=user.id,
        project_id=body.get("project_id") or None,
        title=str(body.get("title") or "Заявка на материалы")[:200],
        status="sent",
        repeated_from=body.get("repeated_from"),
        created_at=now,
    )
    session.add(request)
    for line in lines:
        line.request_id = request.id
        session.add(line)
    for company_id in suppliers:
        session.add(
            RequestRecipient(
                request_id=request.id,
                supplier_company_id=company_id,
                status="sent",
                sent_at=now,
                created_at=now,
                updated_at=now,
            )
        )
    session.flush()
    on_request_sent(request.id, suppliers)
    return request_view(session, user, request.id)


def contact_supplier(session: Session, user, company_id: str) -> dict:
    profile = _profile(session, company_id)
    return create_request(
        session,
        user,
        {
            "title": f"Запрос цены: {profile.name}",
            "supplier_ids": [profile.company_id],
            "lines": [{"name": "Прайс-лист", "unit": "компл", "qty": 1, "category": ""}],
        },
    )


def repeat_request(session: Session, user, request_id: str) -> dict:
    previous = _request_for_buyer(session, user, request_id)
    lines = session.scalars(select(MarketRequestLine).where(MarketRequestLine.request_id == previous.id)).all()
    recipients = session.scalars(select(RequestRecipient).where(RequestRecipient.request_id == previous.id)).all()
    return create_request(
        session,
        user,
        {
            "title": previous.title,
            "project_id": previous.project_id,
            "repeated_from": previous.id,
            "supplier_ids": [row.supplier_company_id for row in recipients],
            "lines": [
                {
                    "catalog_id": line.catalog_id,
                    "name": line.name,
                    "unit": line.unit,
                    "qty": line.qty,
                    "category": line.category,
                }
                for line in lines
            ],
        },
    )


def _request_for_buyer(session: Session, user, request_id: str) -> MarketRequest:
    request = session.get(MarketRequest, request_id)
    if request is None:
        raise MarketError("Заявка не найдена", 404)
    buyer = company_of(session, user, "buyer")
    if request.buyer_company_id != buyer and request.created_by != user.id and user.role != "admin":
        raise MarketError("Недостаточно прав", 403)
    return request


def request_view(session: Session, user, request_id: str, supplier: bool = False) -> dict:
    request = session.get(MarketRequest, request_id)
    if request is None:
        raise MarketError("Заявка не найдена", 404)
    buyer_ok = False
    try:
        _request_for_buyer(session, user, request_id)
        buyer_ok = True
    except MarketError as exc:
        if exc.status != 403:
            raise
    company_id = ""
    if not buyer_ok:
        if user.role not in ("supplier", "admin"):
            raise MarketError("Недостаточно прав", 403)
        company_id = company_of(session, user, "supplier")
        allowed = session.get(RequestRecipient, (request_id, company_id))
        if not allowed and user.role != "admin":
            raise MarketError("Недостаточно прав", 403)
        supplier = True
    lines = session.scalars(select(MarketRequestLine).where(MarketRequestLine.request_id == request_id)).all()
    recipients = session.scalars(select(RequestRecipient).where(RequestRecipient.request_id == request_id)).all()
    offers = session.scalars(select(Offer).where(Offer.request_id == request_id, Offer.active.is_(True))).all()
    if supplier and user.role != "admin":
        offers = [offer for offer in offers if offer.supplier_company_id == company_id]
    else:
        _mark_viewed(session, offers)
    return {
        "id": request.id,
        "title": request.title,
        "status": request.status,
        "created_at": request.created_at,
        "repeated_from": request.repeated_from,
        "project_id": request.project_id,
        "lines": [
            {
                "id": line.id,
                "catalog_id": line.catalog_id,
                "name": line.name,
                "unit": line.unit,
                "qty": line.qty,
                "category": line.category,
            }
            for line in lines
        ],
        "suppliers": [row.supplier_company_id for row in recipients] if buyer_ok else [company_id],
        "offers": [_offer_card(session, offer) for offer in _sorted_offers(session, offers)],
    }


def _mark_viewed(session: Session, offers: list[Offer]) -> None:
    now = _now()
    for offer in offers:
        if offer.status == "sent":
            offer.status = "viewed"
            offer.viewed_at = now


def _sorted_offers(session: Session, offers: list[Offer]) -> list[Offer]:
    def key(offer: Offer) -> tuple:
        card = _offer_card(session, offer)
        return (card["comparable_tiyn"], offer.id)

    return sorted(offers, key=key)


def _offer_card(session: Session, offer: Offer) -> dict:
    lines = session.scalars(select(OfferLine).where(OfferLine.offer_id == offer.id)).all()
    profile = session.get(SupplierProfile, offer.supplier_company_id)
    rating = _rating_row(session, offer.supplier_company_id)
    as_of = _today()
    markets = medians_by_catalog(session, [line.catalog_id for line in lines if line.catalog_id])
    rendered = []
    total = 0
    suspicious = False
    manual = False
    stale_price = False
    analog = False
    comparable_ready = True
    for line in lines:
        incl = incl_tiyn(line.price_tiyn, line.vat_included, line.vat_rate)
        item = session.get(CatalogItem, line.catalog_id) if line.catalog_id else None
        expired = stale(item) if item is not None else False
        judged = _judge_line(session, line.catalog_id, line.name, incl, as_of)
        market = markets.get(line.catalog_id or "", _empty_market())
        overpay_pct, overpay_label = overpay_parts(incl, market.get("median_tiyn"))
        if expired:
            stale_price = True
            comparable_ready = False
        if judged["suspicious"]:
            suspicious = True
        if judged["manual_review"]:
            manual = True
        if line.analog:
            analog = True
        amount = int(round(incl * line.qty))
        if not expired:
            total += amount
        rendered.append(
            {
                "id": line.id,
                "name": line.name,
                "unit": line.unit,
                "qty": line.qty,
                "price_tiyn": line.price_tiyn,
                "vat_label": vat_label(line.vat_included, line.vat_rate),
                "comparable_tiyn": incl,
                "lead_days": line.lead_days,
                "analog": line.analog,
                "analog_label": "аналог" if line.analog else "",
                "stale_label": "цена устарела" if expired else "",
                "suspicious": judged["suspicious"],
                "manual_review": judged["manual_review"],
                "median_tiyn": market.get("median_tiyn"),
                "median_label": market.get("median_label") or LOW_DATA,
                "market_city": market.get("market_city") or "",
                "market_label": market.get("market_label") or "",
                "overpay_pct": overpay_pct,
                "overpay_label": overpay_label,
            }
        )
    header = rendered[0] if len(rendered) == 1 else None
    return {
        "id": offer.id,
        "supplier_id": offer.supplier_company_id,
        "supplier_name": profile.name if profile else offer.supplier_company_id,
        "status": offer.status,
        "version": offer.version,
        "source": offer.source,
        "comparable_tiyn": total,
        "suspicious": suspicious,
        "manual_review": manual,
        "stale": stale_price,
        "analog": analog,
        "analog_label": "аналог" if analog else "",
        "stale_label": "цена устарела" if stale_price else "",
        "vat_label": "без НДС" if profile and (not profile.vat_payer or profile.vat_rate <= 0) else "с НДС",
        "esf_label": "" if profile and profile.esf else "без ЭСФ",
        "best": False,
        "eligible_best": comparable_ready and not stale_price and not suspicious,
        "median_tiyn": header.get("median_tiyn") if header else None,
        "market_city": header.get("market_city") if header else "",
        "market_label": header.get("market_label") if header else "",
        "overpay_pct": header.get("overpay_pct") if header else None,
        "overpay_label": header.get("overpay_label") if header else "",
        "lines": rendered,
        **stars_of(rating, profile),
    }


def comparison(session: Session, user, request_id: str) -> dict:
    view = request_view(session, user, request_id)
    offers = view["offers"]
    eligible = [offer for offer in offers if offer["eligible_best"] and offer["status"] in {"sent", "viewed", "selected"}]
    if eligible:
        best_id = min(eligible, key=lambda offer: (offer["comparable_tiyn"], offer["id"]))["id"]
        for offer in offers:
            offer["best"] = offer["id"] == best_id
    view["offers"] = sorted(offers, key=lambda offer: (offer["comparable_tiyn"], offer["id"]))
    return view


def _recipient_or_raise(session: Session, request_id: str, company_id: str) -> None:
    found = session.get(RequestRecipient, (request_id, company_id))
    if found is None:
        raise MarketError("Заявка не адресована этой компании")


def submit_offer(session: Session, user, body: dict) -> dict:
    profile = _own_profile(session, user)
    request_id = str(body.get("request_id") or "")
    request = session.get(MarketRequest, request_id)
    if request is None:
        raise MarketError("Заявка не найдена", 404)
    _recipient_or_raise(session, request_id, profile.company_id)
    source = _require_source(str(body.get("source") or ""))
    existing = session.scalar(
        select(Offer).where(
            Offer.request_id == request_id,
            Offer.supplier_company_id == profile.company_id,
            Offer.active.is_(True),
        )
    )
    now = _now()
    version = 1
    if existing is not None:
        if existing.status in {"draft", "sent", "viewed"}:
            existing.active = False
            existing.status = "rejected" if existing.status == "draft" else existing.status
            version = existing.version + 1
        else:
            raise MarketError("По этой заявке уже есть предложение", 409)
    offer = Offer(
        id=uuid4().hex[:16],
        request_id=request_id,
        supplier_company_id=profile.company_id,
        status="sent",
        active=True,
        version=version,
        source=source,
        created_at=now,
        sent_at=now,
    )
    session.add(offer)
    raw_lines = list(body.get("lines") or [])
    if not raw_lines:
        raise MarketError("Укажите цены по позициям")
    wrote_observation = False
    request_lines = {
        line.id: line
        for line in session.scalars(select(MarketRequestLine).where(MarketRequestLine.request_id == request_id)).all()
    }
    for raw in raw_lines:
        line_source = _require_source(str(raw.get("source") or source))
        try:
            price = int(raw.get("price_tiyn"))
        except (TypeError, ValueError) as exc:
            raise MarketError("Цена в тиынах, целое число") from exc
        if price <= 0:
            raise MarketError("Цена в тиынах, целое число")
        request_line = request_lines.get(str(raw.get("request_line_id") or ""))
        catalog_id = str(raw.get("catalog_id") or (request_line.catalog_id if request_line else "") or "") or None
        name = str(raw.get("name") or (request_line.name if request_line else "")).strip()
        if not name:
            raise MarketError("Укажите позицию")
        vat_rate = 0 if not profile.vat_payer else int(raw.get("vat_rate") or profile.vat_rate or 16)
        vat_included = False if not profile.vat_payer else bool(raw.get("vat_included", True))
        session.add(
            OfferLine(
                id=uuid4().hex[:16],
                offer_id=offer.id,
                request_line_id=request_line.id if request_line else None,
                catalog_id=catalog_id,
                name=name,
                unit=str(raw.get("unit") or (request_line.unit if request_line else "шт")),
                qty=float(raw.get("qty") or (request_line.qty if request_line else 1)),
                price_tiyn=price,
                vat_included=vat_included,
                vat_rate=vat_rate,
                lead_days=max(0, int(raw.get("lead_days") or 0)),
                analog=bool(raw.get("analog")),
                source=line_source,
            )
        )
        unit = str(raw.get("unit") or (request_line.unit if request_line else "шт"))
        if catalog_id and _append_observation(
            session,
            catalog_id=catalog_id,
            company_id=profile.company_id,
            price_tiyn=price,
            vat_included=vat_included,
            vat_rate=vat_rate,
            source_type="offer",
            source_id=offer.id,
            name=name,
            unit=unit,
        ):
            wrote_observation = True
    try:
        session.flush()
    except IntegrityError as exc:
        raise MarketError("По этой заявке уже есть предложение", 409) from exc
    if wrote_observation:
        recompute_price_index(session, _today())
    return _offer_card(session, offer)


def _position_id(session: Session, catalog_id: str, name: str, unit: str) -> str:
    """Normalized price item when one exists, otherwise the catalog item itself."""

    mapped = _catalog_index(session).get(catalog_id)
    if mapped:
        return mapped
    if session.get(PriceItem, catalog_id) is None:
        session.add(PriceItem(id=catalog_id, name=name[:200] or catalog_id, unit=unit[:20], catalog_ids=[catalog_id]))
        session.flush()
    return catalog_id


def _append_observation(
    session: Session,
    catalog_id: str,
    company_id: str,
    price_tiyn: int,
    vat_included: bool,
    vat_rate: int,
    source_type: str,
    source_id: str,
    name: str,
    unit: str,
) -> bool:
    """Insert one excl-VAT observation. Returns false when there is no source."""

    if source_type not in SOURCE_TYPES or not has_source(source_type, source_id) or price_tiyn <= 0:
        return False
    catalog = session.get(CatalogItem, catalog_id)
    profile = session.get(SupplierProfile, company_id)
    city = (catalog.city if catalog is not None and catalog.city else "") or (profile.city if profile is not None else "")
    item_id = _position_id(session, catalog_id, name, unit)
    position = session.get(PriceItem, item_id)
    stored_rate = vat_rate if vat_rate > 0 else 0
    session.add(
        PriceObservation(
            id=uuid4().hex[:16],
            item_id=item_id,
            catalog_id=catalog_id,
            item_name=position.name if position is not None else name,
            city=city,
            observed_on=_today().isoformat(),
            unit=unit[:20],
            price_tiyn=excl_tiyn(price_tiyn, vat_included, vat_rate),
            vat_included=False,
            vat_rate=stored_rate,
            source_type=source_type,
            source_id=source_id,
            company_id=company_id,
            created_at=_now(),
        )
    )
    return True


def select_offer(session: Session, user, offer_id: str) -> dict:
    """Select one offer and open the deal in a single transaction."""
    offer = session.get(Offer, offer_id)
    if offer is None:
        raise MarketError("Предложение не найдено", 404)
    request = _request_for_buyer(session, user, offer.request_id)
    if session.get_bind().dialect.name == "postgresql":
        session.scalars(select(Offer).where(Offer.request_id == offer.request_id).with_for_update()).all()
        session.refresh(offer)
    existing = session.scalar(select(Deal).where(Deal.offer_id == offer.id))
    if existing is not None:
        return {"deal_id": existing.id, "offer_id": offer.id, "status": "selected"}
    other = session.scalar(select(Deal.id).where(Deal.request_id == offer.request_id))
    if other:
        raise MarketError("Предложение уже выбрано", 409)
    if offer.status in {"rejected", "expired"} or not offer.active:
        raise MarketError("Предложение уже закрыто", 409)
    lines = session.scalars(select(OfferLine).where(OfferLine.offer_id == offer.id)).all()
    lead = max((line.lead_days for line in lines), default=0)
    promised = (_today() + timedelta(days=lead or 1)).isoformat()
    now = _now()
    offer.status = "selected"
    offer.active = True
    for other_offer in session.scalars(
        select(Offer).where(Offer.request_id == offer.request_id, Offer.id != offer.id, Offer.active.is_(True))
    ).all():
        other_offer.active = False
        other_offer.status = "rejected"
    sent = offer.sent_at or offer.created_at
    try:
        sent_dt = datetime.fromisoformat(sent)
        created_dt = datetime.fromisoformat(request.created_at)
        hours = max(0, int((sent_dt - created_dt).total_seconds() // 3600))
    except ValueError:
        hours = 0
    deal = Deal(
        id=uuid4().hex[:16],
        offer_id=offer.id,
        request_id=offer.request_id,
        buyer_company_id=request.buyer_company_id,
        supplier_company_id=offer.supplier_company_id,
        promised_at=promised,
        response_hours=hours,
        created_at=now,
    )
    session.add(deal)
    request.status = "selected"
    session.flush()
    on_offer_selected(request.id, offer.id, request.project_id)
    return {"deal_id": deal.id, "offer_id": offer.id, "status": "selected"}


def accept_deal(session: Session, user, deal_id: str, body: dict) -> dict:
    deal = session.get(Deal, deal_id)
    if deal is None:
        raise MarketError("Сделка не найдена", 404)
    _request_for_buyer(session, user, deal.request_id)
    if deal.accepted_at:
        return {"id": deal.id, "accepted_at": deal.accepted_at}
    today = _today().isoformat()
    deal.delivered_at = str(body.get("delivered_at") or today)[:40]
    deal.accepted_at = today
    deal.on_time = bool(body.get("on_time", deal.delivered_at[:10] <= deal.promised_at[:10]))
    deal.price_matched = bool(body.get("price_matched", True))
    deal.acceptance_ok = bool(body.get("acceptance_ok", True))
    profile = session.get(SupplierProfile, deal.supplier_company_id)
    deal.esf_ok = bool(body.get("esf_ok", bool(profile.esf) if profile else False))
    if _record_invoice(session, deal, body):
        session.flush()
        recompute_price_index(session, _today())
    session.flush()
    recompute_supplier(session, deal.supplier_company_id)
    return {"id": deal.id, "accepted_at": deal.accepted_at}


def _record_invoice(session: Session, deal: Deal, body: dict) -> bool:
    """Write invoice prices in the acceptance transaction, and only with a source."""

    source_id = str(body.get("invoice_source_id") or "").strip()
    if not has_source("invoice", source_id):
        return False
    lines = list(session.scalars(select(OfferLine).where(OfferLine.offer_id == deal.offer_id)).all())
    override = body.get("invoice_price_tiyn")
    wrote = False
    for line in lines:
        if not line.catalog_id:
            continue
        price = line.price_tiyn
        if override is not None and len(lines) == 1:
            try:
                price = int(override)
            except (TypeError, ValueError) as exc:
                raise MarketError("Цена в тиынах, целое число") from exc
        if _append_observation(
            session,
            catalog_id=line.catalog_id,
            company_id=deal.supplier_company_id,
            price_tiyn=price,
            vat_included=line.vat_included,
            vat_rate=line.vat_rate,
            source_type="invoice",
            source_id=source_id,
            name=line.name,
            unit=line.unit,
        ):
            wrote = True
    return wrote


def _scores(body: dict) -> dict[str, int]:
    out = {}
    for key in ("quality", "delivery", "price", "communication"):
        try:
            value = int(body.get(key))
        except (TypeError, ValueError) as exc:
            raise MarketError("Оценка от 1 до 5") from exc
        if value < 1 or value > 5:
            raise MarketError("Оценка от 1 до 5")
        out[key] = value
    return out


def _buyer_deal(session: Session, user, deal: Deal) -> str:
    supplier_member = _member(session, user.id, "supplier")
    if supplier_member is not None and supplier_member.company_id == deal.supplier_company_id:
        raise MarketError("Нельзя оценить свою компанию", 403)
    company_id = company_of(session, user, "buyer")
    if company_id == deal.supplier_company_id:
        raise MarketError("Нельзя оценить свою компанию", 403)
    if company_id != deal.buyer_company_id and user.role != "admin":
        raise MarketError("Оценку ставит покупатель по этой сделке", 403)
    return company_id


def create_review(session: Session, user, deal_id: str, body: dict) -> dict:
    deal = session.get(Deal, deal_id)
    if deal is None:
        raise MarketError("Сделка не найдена", 404)
    if not deal.accepted_at:
        raise MarketError("Оценку можно поставить после приёмки", 403)
    company_id = _buyer_deal(session, user, deal)
    _hit_review_limit(user.id)
    if session.scalar(select(SupplierReview.id).where(SupplierReview.deal_id == deal.id)):
        raise MarketError("Оценка по этой сделке уже есть", 409)
    scores = _scores(body)
    now = _now()
    review = SupplierReview(
        id=uuid4().hex[:16],
        deal_id=deal.id,
        buyer_company_id=company_id,
        supplier_company_id=deal.supplier_company_id,
        author_user_id=user.id,
        comment=str(body.get("comment") or "")[:280],
        supplier_reply=None,
        created_at=now,
        updated_at=now,
        **scores,
    )
    session.add(review)
    try:
        session.flush()
    except IntegrityError as exc:
        raise MarketError("Оценка по этой сделке уже есть", 409) from exc
    recompute_supplier(session, deal.supplier_company_id)
    return _review_dict(review)


def update_review(session: Session, user, review_id: str, body: dict) -> dict:
    review = session.get(SupplierReview, review_id)
    if review is None:
        raise MarketError("Оценка не найдена", 404)
    if review.author_user_id != user.id and user.role != "admin":
        raise MarketError("Недостаточно прав", 403)
    created = datetime.fromisoformat(review.created_at)
    if datetime.now(timezone.utc) - created > timedelta(days=EDIT_DAYS):
        raise MarketError("Правка доступна 14 дней", 403)
    scores = _scores({**_review_dict(review), **body})
    review.quality = scores["quality"]
    review.delivery = scores["delivery"]
    review.price = scores["price"]
    review.communication = scores["communication"]
    if "comment" in body:
        review.comment = str(body.get("comment") or "")[:280]
    review.updated_at = _now()
    recompute_supplier(session, review.supplier_company_id)
    return _review_dict(review)


def reply_review(session: Session, user, review_id: str, text: str) -> dict:
    profile = _own_profile(session, user)
    review = session.get(SupplierReview, review_id)
    if review is None or review.supplier_company_id != profile.company_id:
        raise MarketError("Оценка не найдена", 404)
    reply = " ".join(str(text or "").split())
    if not reply:
        raise MarketError("Напишите ответ")
    review.supplier_reply = reply[:280]
    review.updated_at = _now()
    return _review_dict(review)


def _review_dict(review: SupplierReview) -> dict:
    return {
        "id": review.id,
        "deal_id": review.deal_id,
        "quality": review.quality,
        "delivery": review.delivery,
        "price": review.price,
        "communication": review.communication,
        "comment": review.comment,
        "supplier_reply": review.supplier_reply or "",
        "created_at": review.created_at,
        "updated_at": review.updated_at,
    }


def assigned(session: Session, user) -> list[dict]:
    """Requests addressed to this supplier. The inbox screen belongs to registration."""
    profile = _own_profile(session, user)
    rows = session.scalars(
        select(RequestRecipient).where(RequestRecipient.supplier_company_id == profile.company_id)
    ).all()
    out = []
    for row in rows:
        request = session.get(MarketRequest, row.request_id)
        if request is None or request.id.startswith("seed-"):
            continue
        out.append({"id": request.id, "title": request.title, "status": request.status, "created_at": request.created_at})
    out.sort(key=lambda item: item["created_at"], reverse=True)
    return out
