"""Read and write procurement workbooks. Contacts columns are never returned."""

from __future__ import annotations

import io
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.procurement_calc import strip_contacts

MAX_XLSX_ROWS = 2000
HEADER_SCAN = 10

_BEST = PatternFill("solid", fgColor="BBF7D0")
_EMPTY = PatternFill("solid", fgColor="FDE68A")
_HEADER = PatternFill("solid", fgColor="0F2942")
_HEADER_FONT = Font(color="F8FAFC", bold=True)
_BOLD = Font(bold=True)

_POS = {"№", "n", "no", "номер", "поз", "п п", "п/п", "#"}
_NAME = {"name", "title", "item", "материал", "номенклатура", "название", "наименование", "позиция"}
_UNIT = {"unit", "единица", "ед", "ед изм", "единица измерения"}
_QTY = {"qty", "quantity", "объем", "объём", "количество", "кол-во", "кол во"}
_NOTE = {"note", "comment", "комментарий", "примечание", "прим"}
_SUPPLIER = {"поставщик", "supplier", "организация"}
_PRICE = {"цена", "price", "цена за ед", "цена за единицу"}
_VAT = {"с ндс", "ндс", "vat"}
_AVAILABLE = {"доступно", "наличие", "остаток"}
_DOC = {"документ", "счет", "счёт", "invoice"}
_DATE = {"дата", "date"}
_CONFIRMED = {"подтверждено", "confirmed"}
_REF = {"цена сайта", "ориентир"}
_REF_VAT = {"цена сайта с ндс"}
_DELIVERY = {"срок", "срок поставки", "поставка", "delivery"}


def norm_header(value: object) -> str:
    text = str(value or "").strip().lower().replace("ё", "е")
    text = text.replace(".", " ").replace("_", " ").replace("/", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def is_contact_header(text: str) -> bool:
    if "контакт" in text or "телефон" in text or "phone" in text or "e-mail" in text or "email" in text:
        return True
    return bool(re.search(r"(^|[^a-zа-я])тел\.?($|[^a-zа-я])", text))


def _header_kind(text: str, *, offers: bool) -> str | None:
    if not text or is_contact_header(text):
        return "contact" if text and is_contact_header(text) else None
    if text in _POS or text in {"поз", "поз №"} or text.startswith(("№", "поз ")):
        return "pos"
    if "наименован" in text or text in _NAME:
        return "name"
    if text in _UNIT or text.startswith("ед"):
        return "unit"
    if text in _QTY or "колич" in text or text.startswith("кол"):
        return "qty"
    if "примеч" in text or text in _NOTE:
        return "note"
    if not offers:
        return None
    if text in _SUPPLIER or "поставщик" in text:
        return "supplier"
    if text in _REF_VAT or "цена сайта с ндс" in text:
        return "ref_vat"
    if text in _REF or text.startswith("цена сайта"):
        return "reference"
    if text in _VAT or text == "с ндс":
        return "vat"
    if text in _AVAILABLE or "доступ" in text:
        return "available"
    if text in _DOC:
        return "document"
    if text in _DATE:
        return "date"
    if text in _CONFIRMED or "подтверж" in text:
        return "confirmed"
    if text in _DELIVERY or text.startswith("срок"):
        return "delivery"
    if text in _PRICE or (text.startswith("цена") and "сайт" not in text):
        return "price"
    return None


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def parse_number(value: object) -> tuple[float | None, str | None]:
    if value is None or value == "":
        return None, "missing"
    if isinstance(value, bool):
        return None, "invalid"
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().replace("\u00a0", "").replace(" ", "")
        if not text:
            return None, "missing"
        text = text.replace(",", ".")
        try:
            number = float(text)
        except ValueError:
            return None, "invalid"
    if number <= 0:
        return None, "positive"
    return number, None


def parse_bool(value: object) -> bool | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value == 1:
            return True
        if value == 0:
            return False
        return None
    text = str(value).strip().lower().replace("ё", "е")
    if text in {"да", "yes", "y", "true", "1", "с ндс", "ндс", "включен", "включено", "есть"}:
        return True
    if text in {"нет", "no", "n", "false", "0", "без ндс", "без", "не включен", "не включено"}:
        return False
    return None


def _load(data: bytes):
    if not data.startswith(b"PK"):
        raise ValueError("Нужен файл xlsx")
    try:
        return load_workbook(io.BytesIO(data), data_only=False)
    except Exception as exc:
        raise ValueError("Не удалось прочитать xlsx") from exc


def _detect(ws, *, offers: bool) -> tuple[int, dict[str, int], bool] | None:
    max_row = min(ws.max_row or 1, HEADER_SCAN)
    max_col = min(ws.max_column or 1, 30)
    for index, row in enumerate(ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col), start=1):
        found: dict[str, int] = {}
        contact = False
        for cell in row:
            kind = _header_kind(norm_header(cell.value), offers=offers)
            if kind == "contact":
                contact = True
                continue
            if kind and kind not in found:
                found[kind] = cell.column - 1
        if offers and "price" in found and ("pos" in found or "name" in found):
            return index, found, contact
        if not offers and "name" in found and "qty" in found:
            return index, found, contact
    return None


def _rows(ws, header_row: int, columns: dict[str, int]):
    last_col = max(columns.values()) + 1
    max_row = min(ws.max_row or header_row, header_row + MAX_XLSX_ROWS)
    for excel_row in ws.iter_rows(min_row=header_row + 1, max_row=max_row, max_col=max(last_col, 1)):
        values = [cell.value for cell in excel_row]
        yield excel_row[0].row, values


def _pick(values: list, columns: dict[str, int], key: str) -> object:
    index = columns.get(key)
    if index is None or index >= len(values):
        return None
    return values[index]


def _plain(value: object, limit: int = 500) -> tuple[str, bool]:
    cleaned, removed = strip_contacts(_cell_text(value))
    return cleaned[:limit], removed


def parse_statement(data: bytes) -> dict:
    wb = _load(data)
    detected = None
    for ws in wb.worksheets:
        found = _detect(ws, offers=False)
        if found:
            detected = (ws, found)
            if ws.title.casefold() in {"ведомость", "заявка", "statement"}:
                break
    if not detected:
        raise ValueError("Не найдена строка заголовков")
    ws, (header_row, columns, contact) = detected
    rows = []
    seq = 0
    for excel_row, values in _rows(ws, header_row, columns):
        pos = _cell_text(_pick(values, columns, "pos"))
        raw_title = _cell_text(_pick(values, columns, "name"))
        raw_unit = _cell_text(_pick(values, columns, "unit"))
        raw_note = _cell_text(_pick(values, columns, "note"))
        qty_raw = _pick(values, columns, "qty")
        if not any([pos, raw_title, raw_unit, raw_note, qty_raw not in (None, "")]):
            continue
        title, title_stripped = _plain(raw_title, 300)
        unit, _unit_stripped = _plain(raw_unit, 40)
        note, note_stripped = _plain(raw_note)
        seq += 1
        errors: list[str] = []
        if "pos" in columns:
            pos_no = pos
            if not pos_no:
                errors.append("pos_required")
        else:
            pos_no = str(seq)
        if not title:
            errors.append("name_required")
        if not unit:
            errors.append("unit_required")
        qty, qty_err = parse_number(qty_raw)
        if qty_err in {"missing", "invalid"}:
            errors.append("qty_invalid")
        elif qty_err == "positive":
            errors.append("qty_positive")
        rows.append(
            {
                "row": excel_row,
                "pos_no": pos_no,
                "title": title,
                "unit": unit,
                "qty": qty,
                "note": note,
                "errors": errors,
                "phone_removed": title_stripped or note_stripped,
            }
        )
    return {"rows": rows, "contact_ignored": contact, "header_row": header_row}


def parse_offers(data: bytes) -> dict:
    wb = _load(data)
    detected = None
    for ws in wb.worksheets:
        found = _detect(ws, offers=True)
        if found:
            detected = (ws, found)
            break
    if not detected:
        raise ValueError("Не найдена строка заголовков")
    ws, (header_row, columns, contact) = detected
    rows = []
    for excel_row, values in _rows(ws, header_row, columns):
        pos = _cell_text(_pick(values, columns, "pos"))
        raw_name = _cell_text(_pick(values, columns, "name"))
        supplier = _cell_text(_pick(values, columns, "supplier"))
        price_raw = _pick(values, columns, "price")
        if not any([pos, raw_name, supplier, price_raw not in (None, "")]):
            continue
        title, _stripped = _plain(raw_name, 300)
        supplier_name, supplier_stripped = _plain(supplier, 200)
        price, price_err = parse_number(price_raw)
        available, available_err = parse_number(_pick(values, columns, "available"))
        if available_err:
            available = None
        ref, ref_err = parse_number(_pick(values, columns, "reference"))
        if ref_err:
            ref = None
        errors: list[str] = []
        if "pos" in columns and not pos and not title:
            errors.append("pos_missing")
        if price_err:
            errors.append("price_invalid")
        if "supplier" in columns and (not supplier or not supplier_name):
            errors.append("supplier_required")
        confirmed_raw = _pick(values, columns, "confirmed") if "confirmed" in columns else None
        confirmed = parse_bool(confirmed_raw)
        if "confirmed" not in columns:
            confirmed = True
        elif confirmed is None:
            confirmed = False
        source_label, _source_stripped = _plain(_pick(values, columns, "document"), 200)
        note, _note_stripped = _plain(_pick(values, columns, "note"))
        delivery, _delivery_stripped = _plain(_pick(values, columns, "delivery"), 120)
        rows.append(
            {
                "row": excel_row,
                "pos_no": pos,
                "title": title,
                "supplier_name": supplier_name,
                "price_kzt": price,
                "vat_included": parse_bool(_pick(values, columns, "vat")),
                "qty_available": available,
                "source_label": source_label,
                "offer_date": _cell_text(_pick(values, columns, "date"))[:32],
                "confirmed": bool(confirmed),
                "note": note,
                "delivery_term": delivery,
                "reference_price_kzt": ref,
                "reference_vat_included": parse_bool(_pick(values, columns, "ref_vat")),
                "errors": errors,
                "phone_removed": supplier_stripped,
            }
        )
    return {
        "rows": rows,
        "contact_ignored": contact,
        "supplier_column": "supplier" in columns,
        "header_row": header_row,
    }


def _bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def offers_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Предложения"
    headers = [
        "Поз.",
        "Поставщик",
        "Цена",
        "С НДС",
        "Доступно",
        "Документ",
        "Дата",
        "Подтверждено",
        "Цена сайта",
        "Цена сайта с НДС",
        "Срок",
        "Примечание",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER
        cell.font = _HEADER_FONT
    hint = wb.create_sheet("Подсказка")
    hint["A1"] = "Поз. — номер из ведомости. Поставщик — название организации, без телефона и без ФИО."
    hint["A2"] = "Цена — за единицу. С НДС: да, нет или пусто. Подтверждено: да, если есть счёт или письменный ответ."
    hint["A3"] = "Колонки с контактами не читаются."
    hint.column_dimensions["A"].width = 110
    ws.column_dimensions["B"].width = 28
    return _bytes(wb)


def comparison_workbook(sheet: dict, title: str, savings: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сравнение"
    suppliers = [item["name"] for item in sheet["suppliers"]]
    headers = ["Поз.", "Наименование", "Ед.", "Кол-во", *suppliers, "Статус"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for line in sheet["lines"]:
        values: list[object] = [line["pos_no"], line["title"], line["unit"], line["qty"]]
        for name in suppliers:
            cell_offer = (line.get("cells") or {}).get(name)
            values.append(cell_offer["price_kzt"] if cell_offer else None)
        values.append(line["status"])
        ws.append(values)
        excel_row = ws.max_row
        for offset, name in enumerate(suppliers):
            cell = ws.cell(excel_row, 5 + offset)
            offer = (line.get("cells") or {}).get(name)
            if offer is None:
                cell.fill = _EMPTY
                cell.value = None
                continue
            if offer.get("is_best"):
                cell.fill = _BEST
            if offer.get("is_selected"):
                cell.font = _BOLD
            notes = []
            if offer.get("vat_included") is True:
                notes.append("с НДС")
            elif offer.get("vat_included") is False:
                notes.append("без НДС")
            if offer.get("short"):
                notes.append(f"есть {offer['qty_available']} из {line['qty']}")
            if notes:
                cell.comment = Comment(". ".join(notes), "ANYQ")
        if line["status_code"] == "no_compare":
            ws.cell(excel_row, len(headers)).font = Font(color="64748B")
    total = ["", "Итого", "", ""]
    for item in sheet["suppliers"]:
        total.append(item["total_kzt"])
    total.append("")
    ws.append(total)
    ws.cell(ws.max_row, 2).font = _BOLD
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18 if col > 1 else 12
    ws.column_dimensions["B"].width = 36
    ws.freeze_panes = "B2"
    if savings:
        _savings_sheet(wb, savings)
    wb.properties.title = title[:80]
    return _bytes(wb)


def _savings_sheet(wb: Workbook, savings: dict) -> None:
    ws = wb.create_sheet("Экономия")
    headers = ["Позиция", "Кол-во", "Ед.", "Выбрано", "Второе предложение", "Экономия, ₸", "Причина"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER
        cell.font = _HEADER_FONT
    for row in savings.get("lines") or []:
        selected = row.get("selected") or {}
        second = row.get("second") or {}
        ws.append(
            [
                row.get("title") or "",
                row.get("qty"),
                row.get("unit") or "",
                selected.get("price_kzt"),
                second.get("price_kzt"),
                row.get("saving_kzt") or 0,
                row.get("reason") or "",
            ]
        )
    ws.append([])
    ws.append(["Итого экономия, ₸", savings.get("savings_kzt") or 0])
    ws.append(["База, ₸", savings.get("base_kzt") or 0])
    pct = savings.get("savings_pct") or 0
    ws.append(["Процент", round(float(pct) * 100)])
    for item in savings.get("overprice") or []:
        ws.append(["Выше рынка", item.get("title") or "", item.get("supplier") or "", item.get("total_kzt")])
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22
    ws.column_dimensions["A"].width = 36
