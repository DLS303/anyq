"""Completed-work acts: tenant gate, roles, numbering, exports, and budget fact."""

import io
import os
import sqlite3
import subprocess
import sys
import threading
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.exc import IntegrityError

from app.company_features import project_company_id
from app.ks_schema import KsAct, KsActCounter, VatRate
from app.main import app
from app.pf0_tables import AuditLog
from app.sql_backend import Record
from app.store import store
from app.tax import vat_rate

client = TestClient(app)
PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n"
PDF_AUTHOR = b"%PDF-1.4\n1 0 obj<</Author (Hidden Name)>>endobj\ntrailer<<>>\n%%EOF\n"
PDF_JS = b"%PDF-1.4\n1 0 obj<</JavaScript /JS >>endobj\ntrailer<<>>\n%%EOF\n"


def _auth(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _bin() -> str:
    return f"000000{uuid4().int % 1_000_000:06d}"


def _register(email: str, bin_no: str, phone: str):
    res = client.post(
        "/v1/auth/register",
        json={
            "email": email,
            "password": "secret12",
            "full_name": "Иван Тестов (вымышл.)",
            "phone": phone,
            "company": 'ТОО "Пример Строй" (вымышл.)',
            "company_bin": bin_no,
            "city": "Almaty",
            "role": "buyer",
            "consent_offer_pd": True,
        },
        headers={"User-Agent": "anyq-test"},
    )
    assert res.status_code == 200, res.text
    return res.json()["data"]


def _promote(email: str, role: str):
    user = store.get_user_by_email(email)
    assert user is not None
    updated = user.model_copy(update={"role": role})
    store.users_by_id[user.id] = updated
    store.users[email.lower()] = updated
    return updated


def _enable(token: str):
    listed = client.get("/v1/org/features", headers=_auth(token))
    assert listed.status_code == 200, listed.text
    row = next(item for item in listed.json()["data"] if item["key"] == "ks_acts")
    assert row["enabled"] is False
    saved = client.post(
        "/v1/org/features",
        json={"key": "ks_acts", "enabled": True, "row_version": row["row_version"], "id": row["id"]},
        headers=_auth(token),
    )
    assert saved.status_code == 200, saved.text


def _project(token: str, title: str) -> str:
    created = client.post(
        "/v1/site/projects",
        json={"title": title, "address": "Test street 1"},
        headers=_auth(token),
    )
    assert created.status_code == 200, created.text
    return created.json()["data"]["id"]


def _member(token: str, pid: str, email: str, role: str) -> None:
    added = client.post(
        f"/v1/site/projects/{pid}/members",
        json={"email": email, "role": role},
        headers=_auth(token),
    )
    assert added.status_code == 200, added.text


def _estimate(pid: str, company_id: str, name: str = "Concrete", qty: str = "10.000") -> str:
    line_id = str(uuid4())
    with store._sql.factory() as session:
        session.add(
            Record(
                collection="estimates",
                doc_id=str(uuid4()),
                project_id=pid,
                company_key=company_id,
                seq=1,
                row_version=1,
                search_text="",
                payload={
                    "lines": [
                        {
                            "id": line_id,
                            "name": name,
                            "unit": "m3",
                            "qty": qty,
                            "unit_price_kzt": 1000,
                        }
                    ]
                },
            )
        )
        session.commit()
    return line_id


def _people():
    previous = os.environ.get("ANYQ_DEMO_MODE")
    os.environ["ANYQ_DEMO_MODE"] = "0"
    suffix = uuid4().hex[:8]
    bin_a = _bin()
    bin_b = _bin()
    owner = _register(f"ks-owner-{suffix}@example.com", bin_a, f"+7700000{uuid4().int % 10000:04d}")
    _promote(owner["user"]["email"], "owner")
    foreman = _register(f"ks-fore-{suffix}@example.com", bin_a, f"+7700001{uuid4().int % 10000:04d}")
    _promote(foreman["user"]["email"], "foreman")
    supervisor = _register(f"ks-sup-{suffix}@example.com", bin_a, f"+7700002{uuid4().int % 10000:04d}")
    _promote(supervisor["user"]["email"], "supervisor")
    customer = _register(f"ks-cus-{suffix}@example.com", bin_a, f"+7700003{uuid4().int % 10000:04d}")
    _promote(customer["user"]["email"], "buyer")
    buyer = _register(f"ks-buy-{suffix}@example.com", bin_a, f"+7700004{uuid4().int % 10000:04d}")
    _promote(buyer["user"]["email"], "buyer")
    foreign = _register(f"ks-out-{suffix}@example.com", bin_b, f"+7700005{uuid4().int % 10000:04d}")
    _promote(foreign["user"]["email"], "owner")
    _enable(owner["token"])
    pid = _project(owner["token"], "Example yard")
    _member(owner["token"], pid, foreman["user"]["email"], "foreman")
    _member(owner["token"], pid, supervisor["user"]["email"], "supervisor")
    _member(owner["token"], pid, customer["user"]["email"], "customer")
    _member(owner["token"], pid, buyer["user"]["email"], "procurement")
    company_id = project_company_id(store.get_project(pid))
    _estimate(pid, company_id, name="=SUM(A1)")
    _enable(foreign["token"])
    other = _project(foreign["token"], "Other yard")
    _estimate(other, project_company_id(store.get_project(other)), name="Other work")
    return previous, {
        "owner": owner["token"],
        "foreman": foreman["token"],
        "supervisor": supervisor["token"],
        "customer": customer["token"],
        "procurement": buyer["token"],
        "foreign": foreign["token"],
        "pid": pid,
        "other": other,
        "company_id": company_id,
        "owner_email": owner["user"]["email"],
    }


@pytest.fixture
def yard():
    previous, data = _people()
    try:
        yield data
    finally:
        if previous is None:
            os.environ.pop("ANYQ_DEMO_MODE", None)
        else:
            os.environ["ANYQ_DEMO_MODE"] = previous


def _draft(token: str, pid: str, day: str, **extra):
    body = {
        "id": str(uuid4()),
        "act_date": day,
        "period_from": day,
        "period_to": day,
        "total_tiyn": 999_999_999,
        "vat_tiyn": 1,
    }
    body.update(extra)
    return client.post(f"/v1/site/projects/{pid}/ks-acts", json=body, headers=_auth(token))


def _flow(token: str, pid: str, day: str, qty: str = "2"):
    created = _draft(token, pid, day)
    assert created.status_code == 200, created.text
    act = created.json()["data"]
    assert act["total_tiyn"] == 0
    line = act["lines"][0]
    saved = client.put(
        f"/v1/site/ks-acts/{act['id']}/lines",
        json={
            "row_version": act["row_version"],
            "total_tiyn": 1,
            "lines": [{"id": line["id"], "qty_period": qty, "amount_tiyn": 1}],
        },
        headers=_auth(token),
    )
    assert saved.status_code == 200, saved.text
    act = saved.json()["data"]
    submitted = client.post(
        f"/v1/site/ks-acts/{act['id']}/submit",
        json={"row_version": act["row_version"]},
        headers=_auth(token),
    )
    assert submitted.status_code == 200, submitted.text
    return submitted.json()["data"]


def _sign(token: str, act: dict, raw: bytes = PDF, key: str | None = None):
    headers = _auth(token)
    if key:
        headers["Idempotency-Key"] = key
    return client.post(
        f"/v1/site/ks-acts/{act['id']}/sign",
        files={"file": ("scan.pdf", raw, "application/pdf")},
        data={"row_version": str(act["row_version"])},
        headers=headers,
    )


def test_roles_isolation_and_exports(yard):
    denied = _draft(yard["customer"], yard["pid"], "2026-04-01")
    assert denied.status_code == 403
    assert denied.json()["detail"] == {"code": "forbidden"}
    assert _draft(yard["procurement"], yard["pid"], "2026-04-01").status_code == 403

    created = _draft(yard["foreman"], yard["pid"], "2026-04-02")
    assert created.status_code == 200, created.text
    act = created.json()["data"]
    assert act["company_id"] == yard["company_id"]
    assert act["currency"] == "KZT"
    assert act["number"] == 1
    saved = client.put(
        f"/v1/site/ks-acts/{act['id']}/lines",
        json={"row_version": act["row_version"], "lines": [{"id": act["lines"][0]["id"], "qty_period": "2"}]},
        headers=_auth(yard["foreman"]),
    )
    assert saved.status_code == 200, saved.text
    act = saved.json()["data"]
    assert act["total_tiyn"] == 200_000
    assert act["vat_tiyn"] == 32_000
    assert act["total_with_vat_tiyn"] == 232_000
    blocked = client.post(
        f"/v1/site/ks-acts/{act['id']}/submit",
        json={"row_version": act["row_version"]},
        headers=_auth(yard["foreman"]),
    )
    assert blocked.status_code == 403

    submitted = client.post(
        f"/v1/site/ks-acts/{act['id']}/submit",
        json={"row_version": act["row_version"]},
        headers=_auth(yard["supervisor"]),
    )
    assert submitted.status_code == 200, submitted.text
    act = submitted.json()["data"]
    assert act["status"] == "submitted"
    assert client.post(
        f"/v1/site/ks-acts/{act['id']}/cancel",
        json={"row_version": act["row_version"]},
        headers=_auth(yard["owner"]),
    ).status_code == 409

    stale = client.put(
        f"/v1/site/ks-acts/{act['id']}/lines",
        json={"row_version": 1, "lines": [{"id": act["lines"][0]["id"], "qty_period": "1"}]},
        headers=_auth(yard["foreman"]),
    )
    assert stale.status_code == 409

    signed = _sign(yard["owner"], act, PDF_AUTHOR, key="sign-once")
    assert signed.status_code == 200, signed.text
    signed_act = signed.json()["data"]
    assert signed_act["status"] == "signed"
    assert signed_act["budget_fact_written_at"]
    again = _sign(yard["owner"], act, PDF_AUTHOR, key="sign-once")
    assert again.status_code == 200, again.text
    assert again.json()["data"]["budget_fact_written_at"] == signed_act["budget_fact_written_at"]
    repeat = _sign(yard["owner"], signed_act)
    assert repeat.status_code == 200, repeat.text
    assert repeat.json()["data"]["budget_fact_written_at"] == signed_act["budget_fact_written_at"]
    act = signed_act

    edited = client.put(
        f"/v1/site/ks-acts/{act['id']}/lines",
        json={"row_version": act["row_version"], "lines": [{"id": act["lines"][0]["id"], "qty_period": "3"}]},
        headers=_auth(yard["owner"]),
    )
    assert edited.status_code == 409
    assert edited.json()["detail"]["code"] == "immutable"

    budget = client.get(f"/v1/site/projects/{yard['pid']}/object-budget", headers=_auth(yard["owner"]))
    assert budget.status_code == 200, budget.text
    body = budget.json()["data"]
    gross_kzt = 2320
    assert body["ks_fact_kzt"] == gross_kzt
    assert body["fact_by_category"]["works"] == gross_kzt
    assert body["fact_kzt"] == gross_kzt

    foreign_paths = [
        f"/v1/site/projects/{yard['pid']}/ks-acts",
        f"/v1/site/ks-acts/{act['id']}",
        f"/v1/site/ks-acts/{act['id']}?pid={yard['other']}",
        f"/v1/site/ks-acts/{act['id']}/ks2.pdf",
        f"/v1/site/ks-acts/{act['id']}/ks2.xlsx",
        f"/v1/site/ks-acts/{act['id']}/file",
        f"/v1/site/projects/{yard['pid']}/ks3?period_from=2026-04-01&period_to=2026-04-30",
        f"/v1/site/projects/{yard['pid']}/ks3.pdf?period_from=2026-04-01&period_to=2026-04-30",
        f"/v1/site/projects/{yard['pid']}/ks3.xlsx?period_from=2026-04-01&period_to=2026-04-30",
    ]
    for path in foreign_paths:
        res = client.get(path, headers=_auth(yard["foreign"]))
        assert res.status_code == 404, path
        assert res.json()["detail"] == "Не найдено"
    same = client.get(f"/v1/site/ks-acts/{act['id']}?pid={yard['other']}", headers=_auth(yard["owner"]))
    assert same.status_code == 200, same.text
    assert same.json()["data"]["project_id"] == yard["pid"]
    wrote = client.post(
        f"/v1/site/projects/{yard['pid']}/ks-acts?pid={yard['other']}",
        json={"id": str(uuid4()), "act_date": "2026-04-03", "period_from": "2026-04-03", "period_to": "2026-04-03"},
        headers=_auth(yard["foreign"]),
    )
    assert wrote.status_code == 404

    scan = client.get(f"/v1/site/ks-acts/{act['id']}/file", headers=_auth(yard["owner"]))
    assert scan.status_code == 200, scan.text
    assert scan.headers["x-content-type-options"] == "nosniff"
    assert "attachment" in scan.headers["content-disposition"]
    assert b"/Author" not in scan.content
    assert scan.content.startswith(b"%PDF")

    pdf = client.get(f"/v1/site/ks-acts/{act['id']}/ks2.pdf", headers=_auth(yard["owner"]))
    assert pdf.status_code == 200, pdf.text
    assert pdf.content.startswith(b"%PDF")
    assert b"/Author" not in pdf.content
    assert "attachment" in pdf.headers["content-disposition"]

    sheet = client.get(f"/v1/site/ks-acts/{act['id']}/ks2.xlsx", headers=_auth(yard["owner"]))
    assert sheet.status_code == 200, sheet.text
    book = load_workbook(io.BytesIO(sheet.content))
    values = [cell.value for row in book.active.iter_rows() for cell in row if cell.value is not None]
    assert any(isinstance(value, str) and value.startswith("'=") for value in values)
    assert any(isinstance(value, str) and "тенге" in value for value in values)
    assert all(cell.data_type != "f" for row in book.active.iter_rows() for cell in row)

    js = _draft(yard["owner"], yard["pid"], "2026-04-04")
    assert js.status_code == 200, js.text
    moved = client.post(
        f"/v1/site/ks-acts/{js.json()['data']['id']}/submit",
        json={"row_version": 1},
        headers=_auth(yard["owner"]),
    )
    assert moved.status_code == 200, moved.text
    rejected = _sign(yard["owner"], moved.json()["data"], PDF_JS)
    assert rejected.status_code == 415
    assert rejected.json()["detail"]["code"] == "active_content"


def test_ks3_matches_signed_acts_and_warnings(yard):
    first = _flow(yard["owner"], yard["pid"], "2026-05-01", qty="2")
    signed = _sign(yard["owner"], first)
    assert signed.status_code == 200, signed.text
    first = signed.json()["data"]
    second = _flow(yard["owner"], yard["pid"], "2026-05-02", qty="1")
    signed = _sign(yard["owner"], second)
    assert signed.status_code == 200, signed.text
    second = signed.json()["data"]
    report = client.get(
        f"/v1/site/projects/{yard['pid']}/ks3?period_from=2026-05-01&period_to=2026-05-31",
        headers=_auth(yard["owner"]),
    )
    assert report.status_code == 200, report.text
    data = report.json()["data"]
    assert data["period_tiyn"] == first["total_tiyn"] + second["total_tiyn"]
    assert data["period_vat_tiyn"] == first["vat_tiyn"] + second["vat_tiyn"]
    assert data["period_with_vat_tiyn"] == first["total_with_vat_tiyn"] + second["total_with_vat_tiyn"]
    assert set(data["act_ids"]) == {first["id"], second["id"]}

    over = _draft(yard["owner"], yard["pid"], "2026-05-03")
    assert over.status_code == 200, over.text
    act = over.json()["data"]
    line = act["lines"][0]
    saved = client.put(
        f"/v1/site/ks-acts/{act['id']}/lines",
        json={"row_version": act["row_version"], "lines": [{"id": line["id"], "qty_period": "11"}]},
        headers=_auth(yard["owner"]),
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["data"]["lines"][0]["over_contract"] is True
    amount = saved.json()["data"]["lines"][0]["amount_tiyn"]
    assert amount == 11 * 1000 * 100

    bad = client.post(
        f"/v1/site/projects/{yard['pid']}/ks-acts",
        content=b"{",
        headers={**_auth(yard["owner"]), "Content-Type": "application/json"},
    )
    assert bad.status_code == 422
    huge = _draft(yard["owner"], yard["pid"], "2026-05-04", contract_ref="A" * 500)
    assert huge.status_code == 422
    negative = client.put(
        f"/v1/site/ks-acts/{act['id']}/lines",
        json={"row_version": saved.json()["data"]["row_version"], "lines": [{"id": line["id"], "qty_period": "-1"}]},
        headers=_auth(yard["owner"]),
    )
    assert negative.status_code == 422


def test_reject_cancel_and_gapless_numbers(yard):
    created = _draft(yard["owner"], yard["pid"], "2026-06-01")
    assert created.status_code == 200, created.text
    act = created.json()["data"]
    direct = _sign(yard["owner"], act)
    assert direct.status_code == 409
    submitted = client.post(
        f"/v1/site/ks-acts/{act['id']}/submit",
        json={"row_version": act["row_version"]},
        headers=_auth(yard["owner"]),
    )
    assert submitted.status_code == 200, submitted.text
    rejected = client.post(
        f"/v1/site/ks-acts/{act['id']}/reject",
        json={"row_version": submitted.json()["data"]["row_version"], "reason_code": "qty_mismatch"},
        headers=_auth(yard["supervisor"]),
    )
    assert rejected.status_code == 200, rejected.text
    assert rejected.json()["data"]["status"] == "rejected"
    reopened = client.post(
        f"/v1/site/ks-acts/{act['id']}/draft",
        json={"row_version": rejected.json()["data"]["row_version"]},
        headers=_auth(yard["foreman"]),
    )
    assert reopened.status_code == 200, reopened.text
    assert reopened.json()["data"]["status"] == "draft"
    cancelled = client.post(
        f"/v1/site/ks-acts/{act['id']}/cancel",
        json={"row_version": reopened.json()["data"]["row_version"]},
        headers=_auth(yard["owner"]),
    )
    assert cancelled.status_code == 200, cancelled.text
    assert cancelled.json()["data"]["status"] == "cancelled"

    pid = _project(yard["owner"], "Number yard")
    _estimate(pid, yard["company_id"])
    user = store.get_user_by_email(yard["owner_email"])
    project = store.get_project(pid)
    from app.ks_acts import Access, create_draft

    access = Access(user=user, project=project, company_id=yard["company_id"])
    numbers: list[int] = []
    errors: list[str] = []
    lock = threading.Lock()

    def create(index: int) -> None:
        day = f"2026-07-{index:02d}"
        try:
            act = create_draft(
                access,
                {"id": str(uuid4()), "act_date": day, "period_from": day, "period_to": day},
            )
        except Exception as exc:  # noqa: BLE001
            with lock:
                errors.append(f"{type(exc).__name__}:{exc}")
            return
        with lock:
            numbers.append(int(act["number"]))

    threads = [threading.Thread(target=create, args=(index,)) for index in range(1, 7)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert errors == []
    assert sorted(numbers) == [1, 2, 3, 4, 5, 6]


def test_flag_off_hides_routes_and_bundle():
    previous = os.environ.get("ANYQ_DEMO_MODE")
    os.environ["ANYQ_DEMO_MODE"] = "0"
    try:
        suffix = uuid4().hex[:8]
        owner = _register(f"ks-off-{suffix}@example.com", _bin(), f"+7700006{uuid4().int % 10000:04d}")
        _promote(owner["user"]["email"], "owner")
        pid = _project(owner["token"], "Hidden yard")
        listed = client.get("/v1/org/features", headers=_auth(owner["token"]))
        row = next(item for item in listed.json()["data"] if item["key"] == "ks_acts")
        assert row["enabled"] is False
        hidden = client.get(f"/v1/site/projects/{pid}/ks-acts", headers=_auth(owner["token"]))
        assert hidden.status_code == 404
        assert hidden.json()["detail"] == "Не найдено"
        bundle = client.get(f"/v1/site/projects/{pid}/bundle", headers=_auth(owner["token"]))
        assert bundle.status_code == 200, bundle.text
        assert bundle.json()["data"]["ks_acts"] == []
    finally:
        if previous is None:
            os.environ.pop("ANYQ_DEMO_MODE", None)
        else:
            os.environ["ANYQ_DEMO_MODE"] = previous


def test_legacy_acts_copy_is_idempotent(yard):
    act_id = str(uuid4())
    with store._sql.factory() as session:
        session.add(
            Record(
                collection="acts",
                doc_id=act_id,
                project_id=yard["pid"],
                company_key=yard["company_id"],
                seq=1,
                row_version=1,
                search_text="",
                payload={
                    "act_date": "2026-02-01",
                    "period_from": "2026-02-01",
                    "period_to": "2026-02-28",
                    "status": "draft",
                    "number": 40,
                    "total_tiyn": 100,
                    "vat_tiyn": 16,
                    "vat_rate_bp": 1600,
                    "lines": [
                        {
                            "id": str(uuid4()),
                            "name": "Legacy line",
                            "unit": "m3",
                            "qty_contract": "1.000",
                            "qty_prev": "0.000",
                            "qty_period": "1.000",
                            "unit_price_tiyn": 100,
                            "amount_tiyn": 100,
                        }
                    ],
                },
            )
        )
        session.commit()
    from app.ks_legacy import copy_legacy_acts

    assert copy_legacy_acts(store._sql.engine) == 1
    assert copy_legacy_acts(store._sql.engine) == 0
    with store._sql.factory() as session:
        row = session.get(KsAct, act_id)
        assert row is not None
        assert row.number == 40
        assert row.total_with_vat_tiyn == 116
        counter = session.get(KsActCounter, {"company_id": yard["company_id"], "project_id": yard["pid"]})
        assert counter is not None
        assert counter.next_no > 40


def test_vat_overlap_is_rejected():
    with store._sql.factory() as session:
        session.add(
            VatRate(
                id=str(uuid4()),
                country="KZ",
                rate_bp=1600,
                valid_from="2026-06-01",
                valid_to=None,
                basis="overlap",
                created_at="2026-10-24T00:00:00+00:00",
                updated_at="2026-10-24T00:00:00+00:00",
                created_by="test",
                updated_by="test",
                row_version=1,
            )
        )
        try:
            session.commit()
        except IntegrityError:
            session.rollback()
            return
    try:
        with pytest.raises(ValueError):
            vat_rate("2026-06-15")
    finally:
        with store._sql.factory() as session:
            session.execute(text("DELETE FROM vat_rates WHERE basis = 'overlap'"))
            session.commit()


def test_audit_keeps_numeric_changes(yard):
    created = _draft(yard["owner"], yard["pid"], "2026-08-01", contract_ref="+77000000001")
    assert created.status_code == 200, created.text
    act = created.json()["data"]
    with store._sql.factory() as session:
        rows = session.scalars(select(AuditLog).where(AuditLog.entity_id == act["id"])).all()
    assert rows
    blob = str([row.changes for row in rows])
    assert "act.create" in {row.action for row in rows}
    assert "+77000000001" not in blob
    assert "Иван" not in blob
    saved = client.put(
        f"/v1/site/ks-acts/{act['id']}/lines",
        json={"row_version": act["row_version"], "lines": [{"id": act["lines"][0]["id"], "qty_period": "1.5"}]},
        headers=_auth(yard["owner"]),
    )
    assert saved.status_code == 200, saved.text
    with store._sql.factory() as session:
        updates = session.scalars(
            select(AuditLog).where(AuditLog.entity_id == act["id"], AuditLog.action == "act.line_update")
        ).all()
    assert updates
    assert "qty_period_milli" in str(updates[-1].changes)


def test_alembic_ks_revision_roundtrip(tmp_path):
    from alembic.config import Config
    from alembic.script import ScriptDirectory

    api = os.path.join(os.path.dirname(__file__), "..")
    script = ScriptDirectory.from_config(Config(os.path.join(api, "alembic.ini")))
    assert script.get_heads() == ["20261024_0001_ks_acts"]
    revision = script.get_revision("20261024_0001_ks_acts")
    assert revision is not None
    assert revision.down_revision == "20261019_0006_pf6"

    db = tmp_path / "ks.sqlite"
    env = os.environ.copy()
    env["DATABASE_URL"] = f"sqlite:///{db}"
    env["ANYQ_ENV"] = "development"
    env["JWT_SECRET"] = "test-only-secret-value-with-32-characters"

    def alembic(*args: str) -> str:
        return subprocess.check_output([sys.executable, "-m", "alembic", *args], cwd=api, env=env, text=True)

    alembic("upgrade", "head")
    alembic("downgrade", "-1")
    alembic("upgrade", "head")
    alembic("check")
    heads = [line for line in alembic("heads").splitlines() if "(head)" in line]
    assert heads == ["20261024_0001_ks_acts (head)"]
    connection = sqlite3.connect(db)
    tables = {row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
    assert {"vat_rates", "company_settings", "ks_act_counters", "ks_acts", "ks_act_lines"} <= tables
    rates = connection.execute("SELECT rate_bp, valid_from, valid_to, basis FROM vat_rates ORDER BY valid_from").fetchall()
    assert rates[0][0] == 1200 and rates[0][1] == "2009-01-01" and rates[0][2] == "2025-12-31"
    assert rates[1][0] == 1600 and rates[1][1] == "2026-01-01" and rates[1][2] is None
    assert "214-VIII" in rates[0][3] and "214-VIII" in rates[1][3]
    counter_columns = {row[1] for row in connection.execute("PRAGMA table_info(ks_act_counters)")}
    assert "id" not in counter_columns
    assert {"company_id", "project_id", "next_no", "updated_at"} <= counter_columns
    idem = {row[1] for row in connection.execute("PRAGMA table_info(idempotency_keys)")}
    assert {"company_id", "scope", "status", "updated_at", "expires_at"} <= idem
    period = connection.execute(
        "SELECT sql FROM sqlite_master WHERE type = 'index' AND name = 'ux_ks_acts_period_open'"
    ).fetchone()[0]
    assert "cancelled" in period and "deleted_at" in period
    connection.close()
