"""Savings follow the second confirmed offer on the same VAT basis.

The reference numbers are the anonymised request from the savings formula:
38 121 ₸ is 23% of a 163 769 ₸ base, and the profile-pipe signal is about 59 472 ₸.
Website prices are not part of the saving.
"""

import io
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.authutil import hash_password
from app.main import app
from app.procurement_calc import compute_savings
from app.procurement_models import ProcLine, ProcOffer, ProcRequest
from app.store import store

client = TestClient(app)

# title, unit, qty, offers (supplier, price, vat_included), reference price, reference vat
FIXTURE = [
    ("Труба PPR Ø20", "м", 94, [("Поставщик A", 450, True), ("Поставщик B", 616, True)], None, None),
    ("Клапан балансировочный Ø20", "шт", 10, [("Поставщик C", 6748.80, True), ("Поставщик D", 8118.90, True)], None, None),
    ("Угол PPR 20", "шт", 33, [("Поставщик A", 45, True), ("Поставщик B", 79, True)], None, None),
    ("Тройник PPR 25", "шт", 14, [("Поставщик A", 100, True), ("Поставщик B", 138, True)], None, None),
    ("Муфта PPR 25×20", "шт", 18, [("Поставщик A", 65, True), ("Поставщик B", 79, True)], None, None),
    ("Отвод Ø50", "шт", 22, [("Поставщик A", 195, True), ("Поставщик B", 311, True)], None, None),
    ("Тройник Ø50", "шт", 9, [("Поставщик A", 460, True), ("Поставщик B", 727, True)], None, None),
    ("Отвод Ø110", "шт", 5, [("Поставщик A", 675, True), ("Поставщик B", 1066, True)], None, None),
    ("Труба Ø50", "м", 30, [("Поставщик F", 786, None), ("Поставщик B", 1003, False)], None, None),
    ("Профтруба 40×40×2", "м", 126, [("Поставщик E", 1354.31, False)], 1099, True),
]


def _models(request_id: str = "req-fixture") -> tuple[list[ProcLine], list[ProcOffer]]:
    lines: list[ProcLine] = []
    offers: list[ProcOffer] = []
    for index, (title, unit, qty, quotes, ref, ref_vat) in enumerate(FIXTURE, start=1):
        line_id = f"line-{index}"
        built: list[ProcOffer] = []
        for supplier, price, vat in quotes:
            built.append(
                ProcOffer(
                    id=f"offer-{index}-{supplier[-1]}",
                    line_id=line_id,
                    supplier_name=supplier,
                    price_kzt=price,
                    vat_included=vat,
                    confirmed=True,
                    offer_date="2026-10-16",
                )
            )
        lines.append(
            ProcLine(
                id=line_id,
                request_id=request_id,
                pos_no=str(index),
                title=title,
                unit=unit,
                qty=qty,
                selected_offer_id=built[0].id,
                reference_price_kzt=ref,
                reference_vat_included=ref_vat,
            )
        )
        offers.extend(built)
    return lines, offers


def _register(email: str, company: str = "Строй Тест") -> tuple[str, str]:
    res = client.post(
        "/v1/auth/register",
        json={
            "email": email,
            "password": "secret12",
            "full_name": "Test User",
            "phone": "+77001112233",
            "company": company,
            "company_bin": "123456789012",
            "city": "Almaty",
            "role": "buyer",
            "consent_offer_pd": True,
        },
    )
    assert res.status_code == 200, res.text
    data = res.json()["data"]
    return data["token"], data["user"]["id"]


def _auth(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _owner() -> str:
    store.credentials["u-owner"] = hash_password("owner-pass-savings")
    owner = store.get_user("u-owner")
    assert owner is not None
    owner.company = "Строй Тест"
    store._save()
    res = client.post(
        "/v1/auth/login",
        json={"email": "owner@demo.superapp", "password": "owner-pass-savings", "consent_offer_pd": True},
    )
    assert res.status_code == 200, res.text
    return res.json()["data"]["token"]


def test_reference_request_savings_and_overprice_signal():
    lines, offers = _models()
    result = compute_savings(lines, offers, vat_rate=0.16, overprice_threshold=0.15)
    assert result["savings_kzt"] == 38121
    assert result["base_kzt"] == 163769
    assert round(result["savings_pct"] * 100) == 23
    reasons = [row["reason"] for row in result["lines"]]
    assert any("Труба Ø50" in reason for reason in reasons)
    pipe = next(row for row in result["lines"] if row["title"] == "Труба Ø50")
    assert pipe["included"] is False
    profile = next(item for item in result["overprice"] if item["title"] == "Профтруба 40×40×2")
    assert profile["per_unit"] == pytest.approx(472, abs=1)
    assert profile["total"] == pytest.approx(59472, abs=150)
    assert profile["total_kzt"] == pytest.approx(59472, abs=150)
    # The website price on the profile pipe is a signal only.
    assert all(row["title"] != "Профтруба 40×40×2" or row["included"] is False for row in result["lines"])


def test_not_cheapest_selection_saves_nothing():
    line = ProcLine(
        id="line-dear",
        request_id="req",
        pos_no="1",
        title="Цемент",
        unit="кг",
        qty=4,
        selected_offer_id="offer-dear",
    )
    cheap = ProcOffer(
        id="offer-cheap",
        line_id=line.id,
        supplier_name="Поставщик A",
        price_kzt=100,
        vat_included=True,
        confirmed=True,
    )
    dear = ProcOffer(
        id="offer-dear",
        line_id=line.id,
        supplier_name="Поставщик B",
        price_kzt=180,
        vat_included=True,
        confirmed=True,
    )
    result = compute_savings([line], [cheap, dear], vat_rate=0.16)
    assert result["savings_kzt"] == 0
    assert result["base_kzt"] == 0
    row = result["lines"][0]
    assert row["not_cheapest"] is True
    assert row["saving_kzt"] == 0
    assert "выбрано не самое дешёвое" in row["reason"]


def _plant(project_id: str, user_id: str) -> str:
    request = ProcRequest(
        id=uuid4().hex,
        project_id=project_id,
        number=2,
        title="Заявка 2",
        status="decided",
        created_at="2026-10-19T00:00:00+00:00",
        created_by=user_id,
    )
    lines, offers = _models(request.id)
    with store._lock:
        store.proc_requests[request.id] = request
        for line in lines:
            store.proc_lines[line.id] = line
        for offer in offers:
            store.proc_offers[offer.id] = offer
        store._save_unlocked()
    return request.id


def test_savings_endpoint_matches_formula_and_hides_it_from_outsiders():
    owner = _owner()
    outsider, _outsider_id = _register("savings-outsider@example.com", "Чужая компания")
    same_company, _same_id = _register("savings-same-co@example.com", "Строй Тест")
    created = client.post(
        "/v1/site/projects",
        json={"title": "Объект экономии", "address": "Тестовый адрес"},
        headers=_auth(owner),
    )
    assert created.status_code == 200, created.text
    project_id = created.json()["data"]["id"]
    request_id = _plant(project_id, "u-owner")

    denied = client.get(f"/v1/procurement/requests/{request_id}/savings", headers=_auth(outsider))
    assert denied.status_code == 403
    missing = client.get("/v1/procurement/requests/missing-savings/savings", headers=_auth(owner))
    assert missing.status_code == 404

    saved = client.get(f"/v1/procurement/requests/{request_id}/savings", headers=_auth(owner))
    assert saved.status_code == 200, saved.text
    data = saved.json()["data"]
    assert data["savings_kzt"] == 38121
    assert data["base_kzt"] == 163769
    assert round(data["savings_pct"] * 100) == 23
    assert any("Труба Ø50" in row["reason"] for row in data["lines"])
    profile = next(item for item in data["overprice"] if "Профтруба" in item["title"])
    assert profile["per_unit"] == pytest.approx(472, abs=1)
    assert profile["total"] == pytest.approx(59472, abs=150)

    project = client.get(f"/v1/procurement/projects/{project_id}/savings", headers=_auth(owner))
    assert project.status_code == 200, project.text
    assert project.json()["data"]["savings_kzt"] == 38121
    assert round(project.json()["data"]["savings_pct"] * 100) == 23

    company = client.get("/v1/procurement/company-savings", headers=_auth(owner))
    assert company.status_code == 200, company.text
    body = company.json()["data"]
    row = next(item for item in body["objects"] if item["project_id"] == project_id)
    assert row["savings_kzt"] == 38121
    assert round(row["savings_pct"] * 100) == 23
    assert body["savings_kzt"] == sum(item["savings_kzt"] for item in body["objects"])
    assert body["savings_kzt"] >= 38121

    leaked = client.get("/v1/procurement/company-savings", headers=_auth(same_company))
    assert leaked.status_code == 200
    assert leaked.json()["data"]["objects"] == []
    assert leaked.json()["data"]["savings_kzt"] == 0
    assert client.get(f"/v1/procurement/projects/{project_id}/savings", headers=_auth(same_company)).status_code == 403

    exported = client.get(f"/v1/procurement/requests/{request_id}/export.xlsx", headers=_auth(owner))
    assert exported.status_code == 200
    book = load_workbook(io.BytesIO(exported.content))
    assert book["Экономия"]["B8"].value == 38121 or any(
        cell.value == 38121 for row in book["Экономия"].iter_rows() for cell in row
    )
    assert client.get(f"/v1/procurement/requests/{request_id}/export.xlsx", headers=_auth(outsider)).status_code == 403
