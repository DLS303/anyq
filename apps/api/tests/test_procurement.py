import io
import json

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.authutil import hash_password
from app.main import app
from app.store import store
from app.store_procurement import dump_lists, read_collections

client = TestClient(app)
PHONE = "+7 700 111 22 33"


def _register(email: str) -> str:
    res = client.post(
        "/v1/auth/register",
        json={
            "email": email,
            "password": "secret12",
            "full_name": "Test User",
            "phone": "+77001112233",
            "company": "Test LLP",
            "company_bin": "123456789012",
            "city": "Almaty",
            "role": "buyer",
            "consent_offer_pd": True,
        },
    )
    assert res.status_code == 200, res.text
    return res.json()["data"]["token"]


def _auth(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


def _owner() -> str:
    store.credentials["u-owner"] = hash_password("owner-pass-proc")
    res = client.post(
        "/v1/auth/login",
        json={"email": "owner@demo.superapp", "password": "owner-pass-proc", "consent_offer_pd": True},
    )
    assert res.status_code == 200, res.text
    return res.json()["data"]["token"]


def _xlsx(fill) -> bytes:
    wb = Workbook()
    fill(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _statement() -> bytes:
    def fill(ws):
        ws.title = "Ведомость"
        ws.append(["Ведомость", None, None, None])
        ws.append([None, None, None, None])
        ws.append([None, None, None, None])
        ws.append(["№", "Наименование", "Ед.", "Кол-во", "Примечание", "Контакт"])
        ws.append([1, "Цемент", "кг", 10, "мешки", PHONE])
        ws.append([2, "Песок", "м3", "много", "", ""])
        ws.append([3, "", "шт", 5, "", ""])
        ws.append([4, "Арматура", "т", 0, "", ""])
        ws.append(["9", "Доска", "м3", 3, "", ""])

    return _xlsx(fill)


def _english_bad() -> bytes:
    def fill(ws):
        ws.append(["Name", "Unit", "Quantity", "Note"])
        ws.append(["Brick", "pcs", 4, "ok"])
        ws.append(["", "pcs", "", "no name and no qty"])
        ws.append(["Sand", "", -2, ""])

    return _xlsx(fill)


def test_legacy_db_without_procurement_keys():
    loaded = read_collections({"change_orders": []})
    assert loaded["proc_requests"] == {}
    assert loaded["proc_lines"] == {}
    assert loaded["proc_offers"] == {}
    assert loaded["proc_docs"] == {}


def test_membership_blocks_outsiders():
    outsider = _register("proc-outsider@example.com")
    denied = client.get("/v1/procurement/projects/pr-besagash/requests", headers=_auth(outsider))
    assert denied.status_code == 403
    missing = client.get("/v1/procurement/projects/missing-project/requests", headers=_auth(outsider))
    assert missing.status_code == 404
    no_token = client.get("/v1/procurement/projects/pr-besagash/requests")
    assert no_token.status_code == 401
    template = client.get("/v1/procurement/offers-template.xlsx", headers=_auth(outsider))
    assert template.status_code == 403

    owner = _owner()
    created = client.post(
        "/v1/procurement/projects/pr-besagash/requests",
        json={"title": "Заявка на бетон"},
        headers=_auth(owner),
    )
    assert created.status_code == 200, created.text
    rid = created.json()["data"]["id"]
    assert created.json()["data"]["number"] >= 1

    again = client.post(
        "/v1/procurement/projects/pr-besagash/requests",
        json={"title": "Заявка на бетон"},
        headers={**_auth(owner), "Idempotency-Key": "proc-create-1"},
    )
    replay = client.post(
        "/v1/procurement/projects/pr-besagash/requests",
        json={"title": "Заявка на бетон"},
        headers={**_auth(owner), "Idempotency-Key": "proc-create-1"},
    )
    assert again.status_code == 200
    assert replay.json()["data"]["id"] == again.json()["data"]["id"]

    for path in (
        f"/v1/procurement/requests/{rid}/sheet",
        f"/v1/procurement/requests/{rid}/export.xlsx",
    ):
        assert client.get(path, headers=_auth(outsider)).status_code == 403
    assert client.get("/v1/procurement/requests/missing-request/sheet", headers=_auth(owner)).status_code == 404
    allowed = client.get("/v1/procurement/offers-template.xlsx", headers=_auth(owner))
    assert allowed.status_code == 200
    assert allowed.content.startswith(b"PK")


def test_statement_preview_skips_bad_rows_and_ignores_contacts():
    owner = _owner()
    created = client.post(
        "/v1/procurement/projects/pr-besagash/requests",
        json={"title": "Ведомость материалов"},
        headers=_auth(owner),
    )
    rid = created.json()["data"]["id"]
    preview = client.post(
        f"/v1/procurement/requests/{rid}/import-lines",
        params={"preview": "1"},
        files={"file": ("ved.xlsx", _statement(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=_auth(owner),
    )
    assert preview.status_code == 200, preview.text
    body = preview.json()["data"]
    assert body["preview"] is True
    assert body["created"] == 0
    assert body["contact_ignored"] is True
    assert PHONE not in preview.text
    codes = {(item["row"], code) for item in body["warnings"] for code in (item["code"],)}
    assert (6, "qty_invalid") in {(item["row"], item["code"]) for item in body["warnings"]}
    assert "name_required" in {item["code"] for item in body["warnings"]}
    assert "qty_positive" in {item["code"] for item in body["warnings"]}
    listed = client.get(f"/v1/procurement/requests/{rid}/sheet", headers=_auth(owner))
    assert listed.json()["data"]["summary"]["lines"] == 0

    saved = client.post(
        f"/v1/procurement/requests/{rid}/import-lines",
        files={"file": ("ved.xlsx", _statement(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=_auth(owner),
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["data"]["created"] == 2
    assert saved.json()["data"]["skipped"] == 3
    assert PHONE not in saved.text
    blob = json.dumps(dump_lists(store), ensure_ascii=False)
    assert "700 111 22 33" not in blob
    assert "7001112233" not in blob

    repeated = client.post(
        f"/v1/procurement/requests/{rid}/import-lines",
        files={"file": ("ved.xlsx", _statement(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=_auth(owner),
    )
    assert repeated.json()["data"]["created"] == 0
    assert repeated.json()["data"]["updated"] == 2
    sheet = client.get(f"/v1/procurement/requests/{rid}/sheet", headers=_auth(owner))
    assert sheet.json()["data"]["summary"]["lines"] == 2
    titles = {row["title"] for row in sheet.json()["data"]["lines"]}
    assert titles == {"Цемент", "Доска"}

    english = client.post(
        f"/v1/procurement/requests/{rid}/import-lines",
        params={"preview": "1"},
        files={"file": ("en.xlsx", _english_bad(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=_auth(owner),
    )
    assert english.status_code == 200, english.text
    warning_codes = {item["code"] for item in english.json()["data"]["warnings"]}
    assert "name_required" in warning_codes
    assert "qty_invalid" in warning_codes or "qty_positive" in warning_codes
    assert "unit_required" in warning_codes
    good = [row for row in english.json()["data"]["rows"] if not row["errors"]]
    assert [row["title"] for row in good] == ["Brick"]
    assert codes  # preview of the Russian file reported row errors


def _offer(owner: str, line_id: str, supplier: str, price: float, vat: bool | None, **extra):
    payload = {
        "supplier_name": supplier,
        "price_kzt": price,
        "vat_included": vat,
        "confirmed": True,
        **extra,
    }
    res = client.post(f"/v1/procurement/lines/{line_id}/offers", json=payload, headers=_auth(owner))
    assert res.status_code == 200, res.text
    return res.json()["data"]


def test_sheet_totals_best_price_and_duplicate_invoices():
    owner = _owner()
    outsider = _register("proc-sheet-out@example.com")
    created = client.post(
        "/v1/procurement/projects/pr-besagash/requests",
        json={"title": "Сравнение цен"},
        headers=_auth(owner),
    )
    rid = created.json()["data"]["id"]
    lines_file = _xlsx(
        lambda ws: (
            ws.append(["№", "Наименование", "Ед.", "Кол-во"]),
            ws.append(["1", "Цемент", "кг", 10]),
            ws.append(["2", "Щебень", "м3", 2]),
            ws.append(["3", "Доска", "шт", 5]),
        )
    )
    imported = client.post(
        f"/v1/procurement/requests/{rid}/import-lines",
        files={"file": ("lines.xlsx", lines_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=_auth(owner),
    )
    assert imported.status_code == 200, imported.text
    sheet = client.get(f"/v1/procurement/requests/{rid}/sheet", headers=_auth(owner)).json()["data"]
    by_pos = {row["pos_no"]: row for row in sheet["lines"]}
    cement = _offer(owner, by_pos["1"]["id"], "Поставщик A", 100, True)
    best = _offer(owner, by_pos["1"]["id"], "Поставщик B", 80, True, qty_available=4, delivery_term="5 дней")
    _offer(owner, by_pos["2"]["id"], "Поставщик A", 50, True)
    _offer(owner, by_pos["2"]["id"], "Поставщик B", 40, False)
    bad_offers = _xlsx(
        lambda ws: (
            ws.append(["Поз.", "Поставщик", "Цена", "С НДС", "Подтверждено"]),
            ws.append(["1", "Поставщик C", "нет", "да", "да"]),
            ws.append(["99", "Поставщик C", 10, "да", "да"]),
            ws.append(["3", "Поставщик C", 15, "да", "да"]),
        )
    )
    offer_import = client.post(
        f"/v1/procurement/requests/{rid}/import-offers",
        files={"file": ("offers.xlsx", bad_offers, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=_auth(owner),
    )
    assert offer_import.status_code == 200, offer_import.text
    imported_codes = {item["code"] for item in offer_import.json()["data"]["warnings"]}
    assert "price_invalid" in imported_codes
    assert "pos_missing" in imported_codes
    assert offer_import.json()["data"]["created"] == 1

    sheet = client.get(f"/v1/procurement/requests/{rid}/sheet", headers=_auth(owner))
    assert sheet.status_code == 200, sheet.text
    body = sheet.json()["data"]
    totals = {item["name"]: item["total_kzt"] for item in body["suppliers"]}
    assert totals["Поставщик A"] == 1100
    assert totals["Поставщик B"] == 880
    assert totals["Поставщик C"] == 75
    rows = {row["pos_no"]: row for row in body["lines"]}
    assert rows["1"]["status"] == "сравнено"
    assert rows["1"]["best_offer_id"] == best["id"]
    assert rows["1"]["short_qty"] is True
    assert rows["2"]["status"] == "нет сравнения"
    assert rows["2"]["best_offer_id"] is None
    assert rows["2"]["comparable"] is False
    assert rows["3"]["status"] == "одно предложение"
    assert body["summary"] == {"lines": 3, "closed": 0, "no_price": 0}

    chosen = client.post(f"/v1/procurement/requests/{rid}/select-best", headers=_auth(owner))
    assert chosen.status_code == 200, chosen.text
    chosen_rows = {row["pos_no"]: row for row in chosen.json()["data"]["lines"]}
    assert chosen_rows["1"]["selected_offer_id"] == best["id"]
    assert chosen_rows["2"]["selected_offer_id"] is None
    assert chosen_rows["3"]["selected_offer_id"]
    assert chosen.json()["data"]["summary"]["closed"] == 2
    per_line = client.post(
        f"/v1/procurement/lines/{by_pos['2']['id']}/select",
        json={"offer_id": cement["id"]},
        headers=_auth(owner),
    )
    assert per_line.status_code == 404

    exported = client.get(f"/v1/procurement/requests/{rid}/export.xlsx", headers=_auth(owner))
    assert exported.status_code == 200
    book = load_workbook(io.BytesIO(exported.content))
    ws = book["Сравнение"]
    green = yellow = False
    values = []
    for row in ws.iter_rows():
        for cell in row:
            values.append(cell.value)
            rgb = str(getattr(cell.fill.fgColor, "rgb", "") or "")
            if cell.value == 80 and rgb.endswith("BBF7D0"):
                green = True
            if rgb.endswith("FDE68A"):
                yellow = True
    assert green
    assert yellow
    assert 1100 in values or 1100.0 in values
    assert "Поставщик A" in values

    pdf = b"%PDF-1.4\n1 0 obj<</Type/Catalog>>endobj\ntrailer<<>>\n%%EOF\n"
    other = b"%PDF-1.4\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\ntrailer<<>>\n%%EOF\n"
    fields = {"supplier_name": "Поставщик A", "doc_number": "3011", "doc_date": "2026-10-16", "amount_kzt": "15000"}
    first = client.post(
        f"/v1/procurement/requests/{rid}/documents",
        data=fields,
        files={"file": ("invoice.pdf", pdf, "application/pdf")},
        headers=_auth(owner),
    )
    assert first.status_code == 200, first.text
    assert first.json()["data"]["duplicate"] is False
    second = client.post(
        f"/v1/procurement/requests/{rid}/documents",
        data=fields,
        files={"file": ("invoice.pdf", pdf, "application/pdf")},
        headers=_auth(owner),
    )
    assert second.status_code == 200, second.text
    assert second.json()["data"]["duplicate"] is True
    assert second.json()["data"]["duplicate_reason"] == "file"
    numbered = client.post(
        f"/v1/procurement/requests/{rid}/documents",
        data=fields,
        files={"file": ("other.pdf", other, "application/pdf")},
        headers=_auth(owner),
    )
    assert numbered.status_code == 200, numbered.text
    assert numbered.json()["data"]["duplicate"] is True
    assert numbered.json()["data"]["duplicate_reason"] == "supplier_number"
    amount = client.post(
        f"/v1/procurement/requests/{rid}/documents",
        data={**fields, "doc_number": "3012", "supplier_name": "Поставщик B"},
        files={"file": ("third.pdf", other + b"x", "application/pdf")},
        headers=_auth(owner),
    )
    assert amount.status_code == 200, amount.text
    assert amount.json()["data"]["duplicate"] is True
    assert amount.json()["data"]["duplicate_reason"] == "amount_date"
    docs = client.get(f"/v1/procurement/requests/{rid}/sheet", headers=_auth(owner)).json()["data"]["documents"]
    assert len(docs) == 4

    file_id = first.json()["data"]["document"]["id"]
    stored = first.json()["data"]["document"]["stored_name"]
    own = client.get(f"/v1/procurement/documents/{file_id}/file", headers=_auth(owner))
    assert own.status_code == 200
    assert own.content.startswith(b"%PDF")
    assert client.get(f"/v1/procurement/documents/{file_id}/file", headers=_auth(outsider)).status_code == 403
    leaked = client.get(f"/uploads/{stored}")
    assert leaked.status_code in {401, 404}
    assert b"%PDF" not in leaked.content
    leaked_auth = client.get(f"/uploads/{stored}", headers=_auth(owner))
    assert leaked_auth.status_code == 404
    assert b"%PDF" not in leaked_auth.content


def test_repeat_order_is_prefilled_and_editable():
    owner = _owner()
    outsider = _register("repeat-out@example.com")
    created = client.post(
        "/v1/procurement/projects/pr-besagash/requests",
        json={"title": "Заявка на повтор"},
        headers=_auth(owner),
    )
    assert created.status_code == 200, created.text
    rid = created.json()["data"]["id"]
    lines_file = _xlsx(
        lambda ws: (
            ws.append(["№", "Наименование", "Ед.", "Кол-во"]),
            ws.append(["1", "Цемент", "кг", 10]),
            ws.append(["2", "Доска", "шт", 4]),
        )
    )
    imported = client.post(
        f"/v1/procurement/requests/{rid}/import-lines",
        files={"file": ("lines.xlsx", lines_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=_auth(owner),
    )
    assert imported.status_code == 200, imported.text
    sheet = client.get(f"/v1/procurement/requests/{rid}/sheet", headers=_auth(owner)).json()["data"]
    _offer(owner, sheet["lines"][0]["id"], "Поставщик A", 100, True)
    denied = client.post(f"/v1/procurement/requests/{rid}/repeat", headers=_auth(outsider))
    assert denied.status_code == 403
    assert client.get("/v1/procurement/projects/pr-besagash/picker", headers=_auth(outsider)).status_code == 403

    repeated = client.post(f"/v1/procurement/requests/{rid}/repeat", headers=_auth(owner))
    assert repeated.status_code == 200, repeated.text
    body = repeated.json()["data"]
    assert body["request"]["id"] != rid
    assert body["request"]["project_id"] == "pr-besagash"
    assert body["request"]["repeated_from"] == rid
    assert body["request"]["title"] == "Заявка на повтор"
    copied = [(row["title"], row["qty"], row["unit"], row["selected_offer_id"]) for row in body["lines"]]
    assert copied == [("Цемент", 10.0, "кг", None), ("Доска", 4.0, "шт", None)]
    new_id = body["request"]["id"]
    new_sheet = client.get(f"/v1/procurement/requests/{new_id}/sheet", headers=_auth(owner)).json()["data"]
    assert new_sheet["summary"]["lines"] == 2
    assert new_sheet["documents"] == []
    assert all(row["offers"] == [] for row in new_sheet["lines"])
    cement = next(row for row in body["lines"] if row["title"] == "Цемент")
    patched = client.patch(
        f"/v1/procurement/lines/{cement['id']}",
        json={"qty": 12},
        headers=_auth(owner),
    )
    assert patched.status_code == 200, patched.text
    assert patched.json()["data"]["qty"] == 12
    original = client.get(f"/v1/procurement/requests/{rid}/sheet", headers=_auth(owner)).json()["data"]
    old_cement = next(row for row in original["lines"] if row["title"] == "Цемент")
    assert old_cement["qty"] == 10

    headers = {**_auth(owner), "Idempotency-Key": "repeat-once"}
    first = client.post(f"/v1/procurement/requests/{rid}/repeat", headers=headers)
    second = client.post(f"/v1/procurement/requests/{rid}/repeat", headers=headers)
    assert first.status_code == 200, first.text
    assert second.json()["data"]["request"]["id"] == first.json()["data"]["request"]["id"]

    picker = client.get("/v1/procurement/projects/pr-besagash/picker", params={"q": "цем"}, headers=_auth(owner))
    assert picker.status_code == 200, picker.text
    assert any(item["title"] == "Цемент" and item["source"] == "recent" for item in picker.json()["data"]["recent"])

    quick = client.post(
        "/v1/procurement/projects/pr-besagash/quick",
        json={"title": "Песок", "unit": "м3", "qty": 3},
        headers=_auth(owner),
    )
    assert quick.status_code == 200, quick.text
    assert quick.json()["data"]["request"]["status"] == "collecting"
    assert quick.json()["data"]["request"]["shared"] is True
    assert quick.json()["data"]["line"]["qty"] == 3

    from app.schemas import User

    supplier = User(
        id="u-sup-repeat",
        email="sup-repeat@example.com",
        full_name="Поставщик C",
        role="supplier",
        company="Поставщик C",
        phone="+77001112244",
    )
    store._index_user(supplier)
    store.credentials[supplier.id] = hash_password("secret12")
    login = client.post(
        "/v1/auth/login",
        json={"email": supplier.email, "password": "secret12", "consent_offer_pd": True},
    )
    assert login.status_code == 200, login.text
    supplier_token = login.json()["data"]["token"]
    inbox = client.get("/v1/procurement/inbox", headers=_auth(supplier_token))
    assert inbox.status_code == 200, inbox.text
    ids = {item["id"] for item in inbox.json()["data"]}
    assert quick.json()["data"]["request"]["id"] in ids
    assert new_id in ids
    assert client.get("/v1/procurement/inbox", headers=_auth(outsider)).status_code == 403
    hidden = client.get(f"/v1/procurement/requests/{new_id}/sheet", headers=_auth(supplier_token))
    assert hidden.status_code == 403
    line_id = new_sheet["lines"][0]["id"]
    bid = client.post(
        f"/v1/procurement/lines/{line_id}/offers",
        json={"supplier_name": "Другой", "price_kzt": 70, "vat_included": True, "confirmed": True},
        headers=_auth(supplier_token),
    )
    assert bid.status_code == 200, bid.text
    assert bid.json()["data"]["supplier_name"] == "Поставщик C"
    visible = client.get(f"/v1/procurement/requests/{new_id}/sheet", headers=_auth(owner)).json()["data"]
    assert any(offer["supplier_name"] == "Поставщик C" for row in visible["lines"] for offer in row["offers"])
