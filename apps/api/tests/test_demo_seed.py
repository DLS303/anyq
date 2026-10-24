"""Demo startup seed, demo env checks, and fictional contact data."""

from __future__ import annotations

import asyncio
import inspect
import json
import logging
import multiprocessing
import re
import secrets
import string
from pathlib import Path

import pytest
from fastapi import FastAPI, Response
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app.authutil import load_secret, reset_secret_cache
from app.config import api_docs_closed, resolve_cors_origins, session_cookie_params
from app.demo_password import assert_demo_password_policy
from app.demo_seed import DEMO_SEED_VERSION, demo_placeholder_jpeg, run_demo_seed
from app.import_json import import_file
from app.main import _lifespan, app, root
from app.paths import resolve_stored_file
from app.ratelimit import login_hash_secret
from app.security import SecureCookieMiddleware
from app.sql_backend import guard_json_store_read, read_json_store
from app.store import MemoryStore

FICTIONAL_PHONE = re.compile(r"^\+7 000 000 00 \d{2}$")
LAND_WORD = re.compile(r"участ(?!ник)|земл", re.IGNORECASE)
PHONE = re.compile(r"\+\d[\d\s()\-]{6,}\d")
EMAIL = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
ALLOWED_EMAIL = re.compile(r"^[^@\s]+@(?:demo\.superapp|example\.com)$", re.IGNORECASE)


def _acceptable_demo_password() -> str:
    """A throwaway password that meets the demo rule. Not a stored secret."""
    alphabet = string.ascii_letters + string.digits + "!#$%*+-=?"
    for _ in range(8):
        chars = [
            secrets.choice(string.ascii_lowercase),
            secrets.choice(string.ascii_uppercase),
            secrets.choice(string.digits),
            secrets.choice("!#$%*+-=?"),
        ]
        chars.extend(secrets.choice(alphabet) for _ in range(12))
        secrets.SystemRandom().shuffle(chars)
        password = "".join(chars)
        folded = password.casefold()
        if any(token in folded for token in ("password", "demo", "change-me", "changeme", "change_me")):
            continue
        return password
    raise RuntimeError("could not build a demo password for the test")


def _sqlite(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.delenv("DATABASE_URL", raising=False)
    monkeypatch.setenv("ANYQ_DB_PATH", str(tmp_path / "db.json"))
    monkeypatch.setenv("ANYQ_UPLOADS_DIR", str(tmp_path / "uploads"))
    monkeypatch.setenv("CORS_ORIGINS", "https://anyq-demo.vercel.app")


def _close(*stores: MemoryStore) -> None:
    for item in stores:
        item._sql.engine.dispose()


def _enable_logs(*names: str) -> None:
    """Alembic's fileConfig disables loggers created before it runs."""
    for name in names:
        logger = logging.getLogger(name)
        logger.disabled = False
        logger.propagate = True


def test_seed_contacts_are_fictional():
    root_dir = Path(__file__).resolve().parents[1] / "app"
    paths = sorted(root_dir.glob("seed*.py")) + [root_dir / "demo_seed.py"]
    phones: list[str] = []
    emails: list[str] = []
    for path in paths:
        text = path.read_text(encoding="utf-8")
        phones.extend(PHONE.findall(text))
        emails.extend(EMAIL.findall(text))
    assert phones
    assert emails
    assert [item for item in phones if FICTIONAL_PHONE.fullmatch(item) is None] == []
    assert [item for item in emails if ALLOWED_EMAIL.fullmatch(item) is None] == []


def test_lifespan_calls_demo_seed(monkeypatch):
    monkeypatch.setattr("app.demo_seed.run_demo_seed", lambda target=None: "skipped")

    async def _enter():
        async with _lifespan(app):
            return True

    assert asyncio.run(_enter()) is True


@pytest.mark.parametrize("env_name", ["demo", "development"])
def test_demo_seed_creates_once_and_skips_on_restart(monkeypatch, tmp_path, caplog, env_name):
    _sqlite(monkeypatch, tmp_path)
    password = _acceptable_demo_password()
    monkeypatch.setenv("ANYQ_ENV", env_name)
    monkeypatch.setenv("ANYQ_DEMO_MODE", "1")
    monkeypatch.setenv("ANYQ_DEMO_PASSWORD", password)
    fresh = MemoryStore()
    try:
        _enable_logs("app.demo_seed")
        with caplog.at_level(logging.INFO, logger="app.demo_seed"):
            assert run_demo_seed(fresh) == "created"
        assert f"demo_seed status=created version={DEMO_SEED_VERSION}" in caplog.text
        assert password not in caplog.text
        projects = [fresh.projects[f"demo-prj-almaty-{number}"] for number in (1, 2)]
        assert all("Алматы" in project.address and "(вымышл.)" in project.address for project in projects)
        assert all(project.budget_kzt and project.finish for project in projects)
        for project in projects:
            assert any(line.project_id == project.id for line in fresh.budget_lines.values())
            assert any(task.project_id == project.id for task in fresh.schedule.values())
            assert any(item.project_id == project.id for item in fresh.rfis.values())
            assert any(item.project_id == project.id for item in fresh.proc_requests.values())
            assert project.id in fresh.object_budgets
            log_row = fresh.daily_logs[f"{project.id}-daily"]
            assert log_row.photos
            label = "1" if project.id.endswith("1") else "2"
            stored = resolve_stored_file("public", f"{project.id}/{project.id}-photo-pub.jpg")
            assert stored.read_bytes() == demo_placeholder_jpeg(label)
            assert "DEMO" in inspect.getsource(demo_placeholder_jpeg)
            assert Image.open(stored).format == "JPEG"
        workbook = resolve_stored_file("upload", "demo-prj-almaty-1/demo-estimate.xlsx")
        sheet = load_workbook(workbook).active
        rows = list(sheet.iter_rows(values_only=True))
        assert rows[0][0] == "mark"
        assert {row[0] for row in rows[1:]} == {"DEMO"}
        assert len(rows) == 3
        owner = fresh.get_user("u-owner")
        assert owner is not None
        _user, token = fresh.login("owner@demo.superapp", password, consent=True)
        assert token
        assert password not in token
        snapshot = fresh.projects["demo-prj-almaty-1"].model_dump()
        assert run_demo_seed(fresh) == "skipped"
        assert fresh.projects["demo-prj-almaty-1"].model_dump() == snapshot
    finally:
        _close(fresh)

    restarted = MemoryStore()
    try:
        assert run_demo_seed(restarted) == "skipped"
        assert restarted.projects["demo-prj-almaty-1"].address.endswith("(вымышл.)")
        assert "demo-doc-estimate" in restarted.files
    finally:
        _close(restarted)


def test_demo_seed_without_password_still_seeds_and_login_fails(monkeypatch, tmp_path, caplog):
    _sqlite(monkeypatch, tmp_path)
    monkeypatch.setenv("ANYQ_ENV", "demo")
    monkeypatch.setenv("ANYQ_DEMO_MODE", "1")
    monkeypatch.setenv("ANYQ_DEMO_PASSWORD", "")
    attempt = secrets.token_urlsafe(18)
    _enable_logs("app.store")
    with caplog.at_level(logging.WARNING, logger="app.store"):
        fresh = MemoryStore()
    try:
        assert "ANYQ_DEMO_PASSWORD is unset; demo accounts cannot log in" in caplog.text
        assert "u-owner" not in fresh.credentials
        assert run_demo_seed(fresh) == "created"
        assert "demo-prj-almaty-2" in fresh.projects
        with pytest.raises(ValueError):
            fresh.login("owner@demo.superapp", attempt, consent=True)
        assert attempt not in caplog.text
    finally:
        _close(fresh)


def test_demo_seed_is_absent_in_production(monkeypatch, tmp_path, caplog):
    _sqlite(monkeypatch, tmp_path)
    monkeypatch.setenv("ANYQ_ENV", "production")
    monkeypatch.setenv("ANYQ_DEMO_MODE", "1")
    monkeypatch.setenv("ANYQ_DEMO_PASSWORD", "")
    fresh = MemoryStore()
    try:
        _enable_logs("app.demo_seed")
        with caplog.at_level(logging.INFO, logger="app.demo_seed"):
            assert run_demo_seed(fresh) == "skipped"
        assert f"demo_seed status=skipped version={DEMO_SEED_VERSION}" in caplog.text
        assert "demo-prj-almaty-1" not in fresh.projects
        assert list(fresh.projects) == []
    finally:
        _close(fresh)


@pytest.mark.parametrize(
    ("password", "detail"),
    [
        ("Aa1!short", "at least 16"),
        ("a" * 20, "3 of 4"),
        ("Ab1!passwordxxxx", "obvious"),
        ("Aa1!aaaaDEMOxxxx", "obvious"),
        ("Aa1!aaaachange-me", "obvious"),
    ],
)
def test_demo_password_rejects_short_and_weak(monkeypatch, caplog, password, detail):
    monkeypatch.setenv("ANYQ_ENV", "demo")
    monkeypatch.setenv("ANYQ_DEMO_MODE", "1")
    monkeypatch.setenv("ANYQ_DEMO_PASSWORD", password)
    _enable_logs("app.store", "app.demo_password")
    with caplog.at_level(logging.DEBUG), pytest.raises(RuntimeError) as caught:
        MemoryStore()
    assert "ANYQ_DEMO_PASSWORD" in str(caught.value)
    assert detail in str(caught.value)
    assert password not in str(caught.value)
    assert password not in caplog.text


def test_demo_password_accepts_a_strong_value_without_logging_it(monkeypatch, tmp_path, caplog):
    password = _acceptable_demo_password()
    _sqlite(monkeypatch, tmp_path)
    monkeypatch.setenv("ANYQ_ENV", "demo")
    monkeypatch.setenv("ANYQ_DEMO_MODE", "1")
    monkeypatch.setenv("ANYQ_DEMO_PASSWORD", password)
    _enable_logs("app.store", "app.demo_password")
    with caplog.at_level(logging.DEBUG):
        fresh = MemoryStore()
    try:
        assert password not in caplog.text
        _user, token = fresh.login("owner@demo.superapp", password, consent=True)
        assert token
        assert password not in token
        assert password not in caplog.text
    finally:
        _close(fresh)


def test_password_strength_skips_when_demo_seeding_is_off(monkeypatch):
    monkeypatch.setenv("ANYQ_ENV", "development")
    monkeypatch.setenv("ANYQ_DEMO_MODE", "")
    monkeypatch.setenv("ANYQ_DEMO_PASSWORD", "short")
    assert_demo_password_policy()


def test_production_rejects_demo_password_without_echoing_it(monkeypatch):
    password = secrets.token_urlsafe(24)
    monkeypatch.setenv("ANYQ_ENV", "production")
    monkeypatch.setenv("ANYQ_DEMO_PASSWORD", password)
    with pytest.raises(RuntimeError) as caught:
        MemoryStore()
    assert "ANYQ_DEMO_PASSWORD" in str(caught.value)
    assert password not in str(caught.value)


@pytest.mark.parametrize("env_name", ["demo", "production"])
def test_secure_env_rejects_wildcard_cors_and_missing_secrets(monkeypatch, env_name):
    decoy = "https://not-for-the-error.example"
    monkeypatch.setenv("ANYQ_ENV", env_name)
    monkeypatch.setenv("CORS_ORIGINS", f"{decoy},*")
    with pytest.raises(RuntimeError) as caught:
        resolve_cors_origins()
    assert "CORS_ORIGINS" in str(caught.value)
    assert decoy not in str(caught.value)

    monkeypatch.setenv("CORS_ORIGINS", "")
    with pytest.raises(RuntimeError, match="CORS_ORIGINS"):
        resolve_cors_origins()

    monkeypatch.delenv("JWT_SECRET", raising=False)
    monkeypatch.delenv("SUPERAPP_SECRET", raising=False)
    reset_secret_cache()
    try:
        with pytest.raises(RuntimeError, match="JWT_SECRET"):
            load_secret()
    finally:
        reset_secret_cache()

    monkeypatch.delenv("LOGIN_HASH_SECRET", raising=False)
    with pytest.raises(RuntimeError, match="LOGIN_HASH_SECRET"):
        login_hash_secret()

    assert session_cookie_params()["secure"] is True
    assert api_docs_closed() is True
    assert "docs" not in root()


def test_development_keeps_docs_and_optional_cors(monkeypatch):
    monkeypatch.setenv("ANYQ_ENV", "development")
    monkeypatch.delenv("CORS_ORIGINS", raising=False)
    assert "http://localhost:3000" in resolve_cors_origins()
    assert api_docs_closed() is False
    assert root()["docs"] == "/docs"
    assert session_cookie_params()["secure"] is False
    assert app.docs_url == "/docs"


def test_secure_cookie_middleware_stamps_demo_and_production(monkeypatch):
    mini = FastAPI()
    mini.add_middleware(SecureCookieMiddleware)

    @mini.get("/cookie")
    def cookie():
        response = Response("ok")
        response.set_cookie("sid", "value", httponly=True, samesite="lax")
        return response

    monkeypatch.setenv("ANYQ_ENV", "demo")
    with TestClient(mini) as client:
        stamped = client.get("/cookie")
    assert "secure" in stamped.headers["set-cookie"].lower()

    monkeypatch.setenv("ANYQ_ENV", "development")
    with TestClient(mini) as client:
        plain = client.get("/cookie")
    assert "secure" not in plain.headers["set-cookie"].lower()


@pytest.mark.parametrize("env_name", ["demo", "production"])
def test_json_store_read_raises(monkeypatch, tmp_path, env_name):
    path = tmp_path / "db.json"
    path.write_text(
        json.dumps(
            {
                "users": [
                    {
                        "id": "u-from-json",
                        "email": "json@example.com",
                        "full_name": "JSON",
                        "role": "buyer",
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setenv("ANYQ_ENV", env_name)
    monkeypatch.setenv("ANYQ_DEMO_MODE", "")
    monkeypatch.setenv("ANYQ_DEMO_PASSWORD", "")
    monkeypatch.delenv("DATABASE_URL", raising=False)
    monkeypatch.setenv("ANYQ_DB_PATH", str(path))
    with pytest.raises(RuntimeError) as caught:
        read_json_store(path)
    assert "db.json" in str(caught.value)
    assert "u-from-json" not in str(caught.value)
    with pytest.raises(RuntimeError, match="db.json"):
        guard_json_store_read()
    with pytest.raises(RuntimeError, match="db.json"):
        import_file(path)
    fresh = MemoryStore()
    try:
        assert "u-from-json" not in fresh.users_by_id
        with pytest.raises(RuntimeError, match="db.json"):
            fresh._load_json()
    finally:
        _close(fresh)

    monkeypatch.setenv("ANYQ_ENV", "development")
    assert read_json_store(path)["users"][0]["id"] == "u-from-json"


def _seed_ids(store) -> dict[str, list]:
    photos = [
        (item.original_id, item.public_id)
        for log in store.daily_logs.values()
        for item in log.photos
    ]
    return {
        "projects": sorted(store.projects),
        "tasks": sorted(store.schedule),
        "rfis": sorted(store.rfis),
        "requests": sorted(store.proc_requests),
        "lines": sorted(store.proc_lines),
        "files": sorted(store.files),
        "logs": sorted(store.daily_logs),
        "photos": sorted(photos),
    }


def _land_texts(value, found: list[str]) -> None:
    if isinstance(value, str):
        if LAND_WORD.search(value):
            found.append(value)
        return
    if isinstance(value, dict):
        for item in value.values():
            _land_texts(item, found)
        return
    if isinstance(value, list):
        for item in value:
            _land_texts(item, found)
        return
    dump = getattr(value, "model_dump", None)
    if dump is not None:
        _land_texts(dump(), found)


def test_demo_seeds_construction_without_plots_or_db_json(monkeypatch, tmp_path):
    path = tmp_path / "db.json"
    path.write_text(
        json.dumps(
            {
                "users": [
                    {
                        "id": "u-from-json",
                        "email": "json@example.com",
                        "full_name": "JSON",
                        "role": "buyer",
                    }
                ],
                "plots": [{"id": "p-from-json"}],
            }
        ),
        encoding="utf-8",
    )
    _sqlite(monkeypatch, tmp_path)
    monkeypatch.setenv("ANYQ_DB_PATH", str(path))
    monkeypatch.setenv("ANYQ_ENV", "demo")
    monkeypatch.setenv("ANYQ_DEMO_MODE", "1")
    monkeypatch.setenv("ANYQ_DEMO_PASSWORD", "")
    fresh = MemoryStore()
    try:
        assert list(fresh.plots) == []
        assert "p-from-json" not in fresh.plots
        assert "p-besagash" not in fresh.plots
        assert "u-from-json" not in fresh.users_by_id
        assert list(fresh.reviews) == []
        assert "sp-arch" in fresh.specialists
        assert "ord-topo" not in fresh.orders
        assert "deal-land-1" not in fresh.deals
        assert "deal-design-1" not in fresh.deals
        assert "ord-arch" in fresh.orders
        assert fresh.orders["ord-arch"].plot_id is None
        assert "deal-smr-1" in fresh.deals
        assert fresh.deals["deal-smr-1"].plot_id is None
        assert fresh.deals["deal-design-2"].plot_id is None
        assert "th-topo" not in fresh.threads
        assert "m1" not in fresh.messages
        assert "th-arch" in fresh.threads
        assert fresh.threads["th-arch"].plot_id is None
        assert "m3" in fresh.messages
        assert run_demo_seed(fresh) == "created"
        assert "demo-prj-almaty-1-task-1" in fresh.schedule
        assert "demo-doc-estimate" in fresh.files
        found: list[str] = []
        for name in (
            "specialists",
            "orders",
            "deals",
            "threads",
            "messages",
            "projects",
            "schedule",
            "budget_lines",
            "rfis",
            "leads",
            "crm_contacts",
            "crm_activities",
            "crm_invoices",
            "timesheets",
            "daily_logs",
            "files",
            "proc_requests",
        ):
            _land_texts(getattr(fresh, name), found)
        assert found == []
        for order in fresh.orders.values():
            assert order.plot_id is None
        for deal in fresh.deals.values():
            assert deal.plot_id is None
    finally:
        _close(fresh)

    again = MemoryStore()
    try:
        assert list(again.plots) == []
        assert "sp-arch" in again.specialists
        assert "ord-arch" in again.orders
        assert again.orders["ord-arch"].plot_id is None
        assert "deal-land-1" not in again.deals
        assert "deal-design-2" in again.deals
        assert again.deals["deal-design-2"].plot_id is None
        assert "th-arch" in again.threads
        assert "m3" in again.messages
        assert "u-from-json" not in again.users_by_id
    finally:
        _close(again)


def test_demo_restart_keeps_participation_records(monkeypatch, tmp_path):
    from app.schemas import Deal, Lead, ServiceOrder

    _sqlite(monkeypatch, tmp_path)
    monkeypatch.setenv("ANYQ_ENV", "demo")
    monkeypatch.setenv("ANYQ_DEMO_MODE", "1")
    monkeypatch.setenv("ANYQ_DEMO_PASSWORD", "")
    fresh = MemoryStore()
    try:
        fresh.orders["viewer-ord"] = ServiceOrder(
            id="viewer-ord",
            specialist_id="sp-arch",
            client_id="u-buyer",
            message="Просим подтвердить участие",
            created_at="2026-10-21T00:00:00+00:00",
        )
        fresh.deals["viewer-deal"] = Deal(
            id="viewer-deal",
            funnel="construction",
            stage="request",
            title="Подтверждение участия",
            created_at="2026-10-21T00:00:00+00:00",
        )
        fresh.leads["viewer-lead"] = Lead(
            id="viewer-lead",
            title="Заявка на участие",
            name="Зритель",
            phone="+7 000 000 00 99",
            created_at="2026-10-21T00:00:00+00:00",
        )
        fresh._save()
        before = {
            "orders": sorted(fresh.orders),
            "deals": sorted(fresh.deals),
            "leads": sorted(fresh.leads),
        }
    finally:
        _close(fresh)

    again = MemoryStore()
    try:
        assert sorted(again.orders) == before["orders"]
        assert sorted(again.deals) == before["deals"]
        assert sorted(again.leads) == before["leads"]
        assert again.orders["viewer-ord"].message == "Просим подтвердить участие"
        assert again.deals["viewer-deal"].title == "Подтверждение участия"
        assert again.leads["viewer-lead"].title == "Заявка на участие"
    finally:
        _close(again)


def test_seed_without_marker_does_not_duplicate(monkeypatch, tmp_path):
    from app.sql_backend import Meta

    _sqlite(monkeypatch, tmp_path)
    monkeypatch.setenv("ANYQ_ENV", "demo")
    monkeypatch.setenv("ANYQ_DEMO_MODE", "1")
    monkeypatch.setenv("ANYQ_DEMO_PASSWORD", "")
    fresh = MemoryStore()
    try:
        assert run_demo_seed(fresh) == "created"
        before = _seed_ids(fresh)
        assert "demo-prj-almaty-1-task-1" in before["tasks"]
        assert "demo-prj-almaty-1-rfi" in before["rfis"]
        assert "demo-prj-almaty-1-req" in before["requests"]
        assert "demo-prj-almaty-1-photo-orig" in before["files"]
        assert "demo-doc-estimate" in before["files"]
        with fresh._sql.factory() as session:
            row = session.get(Meta, "demo_seed_version")
            assert row is not None
            session.delete(row)
            session.commit()
        assert run_demo_seed(fresh) == "created"
        assert _seed_ids(fresh) == before
    finally:
        _close(fresh)


def _seed_in_process(barrier, queue) -> None:
    barrier.wait(60)
    from app.demo_seed import run_demo_seed
    from app.store import store

    try:
        queue.put(run_demo_seed(store))
    finally:
        store._sql.engine.dispose()


def test_two_processes_create_the_seed_once(monkeypatch, tmp_path):
    _sqlite(monkeypatch, tmp_path)
    monkeypatch.setenv("ANYQ_ENV", "demo")
    monkeypatch.setenv("ANYQ_DEMO_MODE", "1")
    monkeypatch.setenv("ANYQ_DEMO_PASSWORD", "")
    ctx = multiprocessing.get_context("spawn")
    barrier = ctx.Barrier(2)
    queue = ctx.Queue()
    procs = [ctx.Process(target=_seed_in_process, args=(barrier, queue)) for _ in range(2)]
    for proc in procs:
        proc.start()
    statuses = []
    try:
        for proc in procs:
            proc.join(90)
            if proc.is_alive():
                proc.terminate()
                proc.join(5)
        statuses = [queue.get(timeout=5) for _ in procs]
    finally:
        for proc in procs:
            if proc.is_alive():
                proc.kill()
                proc.join(5)
    assert statuses.count("created") == 1
    assert statuses.count("skipped") == 1
    assert all(proc.exitcode == 0 for proc in procs)

    fresh = MemoryStore()
    try:
        assert sorted(item for item in fresh.projects if item.startswith("demo-prj-")) == [
            "demo-prj-almaty-1",
            "demo-prj-almaty-2",
        ]
        assert list(fresh.files).count("demo-doc-estimate") == 1
    finally:
        _close(fresh)
